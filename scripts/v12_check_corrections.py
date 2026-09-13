"""Compare corrected data cells with the previous published V12 and audit groups."""
from pathlib import Path
import io,json,math,subprocess,zipfile
from datetime import datetime
import openpyxl
from docx import Document
from docx.oxml.ns import qn

root=Path(__file__).resolve().parents[1]
report={'status':'PASS','workbooks':{}}
for key in ['1','2','3','4-2','4-3']:
    name=f'results/V12/附件5/result{key}.xlsx'
    previous=subprocess.check_output(['git','show','e36b89a:'+name],cwd=root)
    old=openpyxl.load_workbook(io.BytesIO(previous),read_only=True,data_only=True)
    now=openpyxl.load_workbook(root/name,data_only=True)
    assert now.sheetnames==old.sheetnames
    count=0
    for ws in now:
        assert all(c.fill.patternType is None and c.font.color.rgb=='00000000'
                   for row in ws for c in row)
        assert (ws.max_row,ws.max_column)==(old[ws.title].max_row,old[ws.title].max_column)
        for a,b in zip(ws.values,old[ws.title].values):
            for x,y in zip(a,b):
                count+=1
                if isinstance(y,(int,float)):
                    assert isinstance(x,(int,float)) and math.isclose(x,y,rel_tol=1e-13,abs_tol=1e-9)
                else:assert x==y,(key,ws.title,x,y)
        if ws.title in ['充放电量','紧急购电量'] and key!='1':
            starts=[r for r in range(2,ws.max_row+1) if isinstance(ws.cell(r,1).value,datetime)]
            assert len(starts)==334
            for i,r in enumerate(starts):
                end=starts[i+1]-1 if i+1<len(starts) else ws.max_row
                assert ws.cell(r,1).border.top.style=='thin'
                assert ws.cell(end,1).border.bottom.style=='thin'
                if ws.title=='充放电量':
                    assert end-r==5
                    assert ws.cell(r,5).value=='00:00' and ws.cell(r+1,5).value=='24:00'
                    assert all(isinstance(ws.cell(t,c).value,(int,float)) for t in range(r,end+1) for c in [3,4])
                for t in range(r,end):assert ws.cell(t,1).border.bottom.style is None
    report['workbooks'][key]={'cells_identical':count,'groups_checked':True}
    now.close();old.close()
old=Document(io.BytesIO(subprocess.check_output(['git','show','e36b89a:reports/完整论文_V12.docx'],cwd=root)))
new=Document(root/'reports/完整论文_V12.docx')
assert [[ [c.text for c in r.cells] for r in t.rows] for t in old.tables]==[
    [[c.text for c in r.cells] for r in t.rows] for t in new.tables]
report['all_61_paper_tables_including_appendix_data_unchanged']=True
changed=[]
for i,(a,b) in enumerate(zip(old.inline_shapes,new.inline_shapes),1):
    def blob(doc,shape):
        rid=shape._inline.find('.//'+qn('a:blip')).get(qn('r:embed'))
        return doc.part.related_parts[rid].blob
    if blob(old,a)!=blob(new,b):changed.append(i)
    assert a.width==b.width and a.height==b.height
assert changed==[11],changed
report['changed_figure_media']=changed
report['all_figure_dimensions_preserved']=True
(root/'artifacts/v12/correction_audit.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf8')
print(json.dumps(report,ensure_ascii=False,indent=2))
