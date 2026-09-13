"""Build V10 from V9 without changing equations, tables, or the abstract."""
from pathlib import Path
import re,json,hashlib,shutil
from docx import Document
from docx.shared import Cm,Pt
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from PIL import Image
from src.v9_document import maths,tables
ROOT=Path(__file__).resolve().parents[1]
CHANGES={
'图10左侧按启用的预报组合列出全年节省，固定50%风险余量与基础执行器，分析信息更新的经济价值；右侧给出基础滚动、负载修正及库存价值控制的费用。两部分分别对应预报选择与执行策略的递进。':
'每次预报发布后，以当前储电量和已知信息优化剩余时域，只执行至下一次更新；已执行部分不再修改。最终有效计划始终与午夜原合同净比较，避免把同一电量的中间修订重复计费。固定50%风险余量与基础执行器时，全部更新相对仅午夜预报全期节省146.92万元；基础滚动、负载修正和库存价值控制的费用依次为1,357.60、1,348.91和1,347.93万元。',
'图 10 预报更新组合的全年收益与第三问最终费用':'图 10 剩余时域决策、状态反馈与合同净结算结构',
'因此，价格变化同时影响计划购电时段、日内调单与当前库存保留。图13按题目指定的四个日期分列，对齐价格、储电量及累计紧急购电的日内轨迹，展示两套策略在不同供需条件下的响应。':
'因此，价格变化同时影响计划购电时段、日内调单与当前库存保留。图13汇总2—12月滚动策略的实际电价与有符号储能功率，展示完整运行期的调度分布。四个指定日期的价格预测见图12，购电与储能结果见附录A。',
'图13第一行展示实际价格与午夜预测，第二行展示日前与滚动策略的储电量，第三行展示当日累计紧急购电量。每行共用纵轴尺度，竖虚线为6、12、18点。6月21日两策略均未发生实质紧急购电；9月23日滚动策略补购较多，说明单日效果不必与全年费用排序一致。':
'图13以0.05元/kWh和0.25 MW划分二维网格，颜色表示各格占滚动策略活跃时段的比例，采用对数色标，不作平滑。全期48,096个区间中，绝对功率大于0.01 kW的活跃区间为42,989个，其余数值零值不纳入分布。正功率表示充电，负功率表示放电。横轴使用事后实际价格作运行描述，不表示调度只由价格决定，也不把未来实际价格用于决策。',
'图 13 波动电价下日前与滚动策略的运行过程':'图 13 滚动策略的实际电价与有符号储能功率分布',
'图14上、下两行分别为日前与滚动策略；左列为月度节省，右列为全期节省及95%配对区块重采样区间，均以千元为单位，各行采用独立且明确标注的线性刻度。采用2,000次、14日区块，日前区间为[59,151.20，186,368.40]元，滚动为[1,953.22，10,500.59]元。区间反映同年固定轨迹的条件波动，不校正模型开发与选择偏差。':
'采用2,000次、14日区块的配对重采样，日前节省的95%区间为[59,151.20，186,368.40]元，滚动为[1,953.22，10,500.59]元。区间反映同年固定轨迹的条件波动，不校正模型开发与选择偏差。',
}
def transform(text):
 for old,new in CHANGES.items():text=text.replace(old,new)
 text=text.replace('图14进一步检验库存价值控制的贡献。','进一步检验库存价值控制的贡献。')
 text=re.sub(r'图(\s*)(1[1-4])(?=\D|$)',lambda m:'图'+m[1]+str(int(m[2])-1),text)
 return text
def main():
 out=ROOT/'artifacts/v10';out.mkdir(exist_ok=True)
 source=ROOT/'reports/完整论文_V9.docx';base=Document(source);d=Document(source)
 md=(source.with_suffix('.md')).read_text(encoding='utf8')
 for old in CHANGES:assert md.count(old)==1,old
 for p in d.paragraphs:
  new=transform(p.text)
  if new!=p.text:
   assert not p._element.findall('.//'+qn('m:oMath'))
   first=p.runs[0];first.text=new
   for r in p.runs[1:]:r.text=''
 names=[Path(p).stem for p in re.findall(r'!\[[^]]*\]\(([^)]+)\)',md)]
 manifest=[]
 removed=[]
 for shape,name in zip(list(d.inline_shapes),names):
  if name in ['q3_updates','inventory_savings']:
   paragraph=shape._inline.getparent()
   while paragraph.tag!=qn('w:p'):paragraph=paragraph.getparent()
   caption=paragraph.getnext()
   removed=[list(d._element.body).index(paragraph),list(d._element.body).index(caption)]
   paragraph.getparent().remove(paragraph);caption.getparent().remove(caption)
   continue
  path=ROOT/f'figures/v10/{name}.png'
  with Image.open(path) as im:w,h=im.size
  shape.width=Cm(15.5);shape.height=Cm(15.5*h/w)
  blip=shape._inline.find('.//'+qn('a:blip'));rid=blip.get(qn('r:embed'))
  d.part.related_parts[rid]._blob=path.read_bytes()
  for old in list(blip):
   if old.tag==qn('a:extLst'):blip.remove(old)
  manifest.append(dict(name=name,width_cm=15.5,height_cm=15.5*h/w,pixels=[w,h],sha256=hashlib.sha256(path.read_bytes()).hexdigest()))
 assert len(d.inline_shapes)==12
 assert tables(d)==tables(base) and maths(d)==maths(base)
 used={node.get(qn('r:embed')) for node in d._element.iter() if node.get(qn('r:embed'))}
 for rid,rel in list(d.part.rels.items()):
  if rel.reltype==RT.IMAGE and rid not in used:del d.part.rels[rid]
 dest=ROOT/'reports/完整论文_V10.docx';d.save(dest)
 md=re.sub(r'!\[[^\]]*\]\([^)]*/q3_updates.png\)\n\n图 10 [^\n]+\n\n','',md)
 md=re.sub(r'!\[[^\]]*\]\([^)]*/inventory_savings.png\)\n\n图 14 [^\n]+\n\n','',md)
 md=transform(md).replace('../figures/v9/','../figures/v10/')
 dest.with_suffix('.md').write_text(md,encoding='utf8')
 shutil.copyfile(ROOT/'reports/完整论文_V9公式.tex',ROOT/'reports/完整论文_V10公式.tex')
 for p in (ROOT/'artifacts/v8').glob('result*.xlsx'):shutil.copyfile(p,out/p.name)
 (out/'figure_manifest.json').write_text(json.dumps(dict(source_data='artifacts/v8/figure_data.mat',source_sha256=hashlib.sha256((ROOT/'artifacts/v8/figure_data.mat').read_bytes()).hexdigest(),backend='MATLAB',figures=manifest),ensure_ascii=False,indent=2),encoding='utf8')
 print('Built V10: 12 figures; 61 unchanged tables; 52 unchanged native equations.')
if __name__=='__main__':main()
