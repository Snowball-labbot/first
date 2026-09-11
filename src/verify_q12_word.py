"""Verify final Word contents against frozen evidence and its rendered pages."""
import argparse
from pathlib import Path
import hashlib
import json
import re
import sys
import zipfile
from collections import Counter
from docx import Document
from docx.oxml.ns import qn
import fitz


def sha(p):
    p=Path(p);data=p.read_bytes()
    if p.suffix in {'.py','.md','.tex','.json'}:data=data.replace(b'\r\n',b'\n')
    return hashlib.sha256(data).hexdigest()


def main():
    parser=argparse.ArgumentParser()
    parser.add_argument('--render-dir',type=Path,default=Path('.qa/word_revision/render_final'))
    args=parser.parse_args()
    root=Path(__file__).resolve().parents[1]
    out=root/'artifacts/q12_revision'
    manifest=json.loads((out/'word_build.json').read_text(encoding='utf-8'))
    expected=json.loads((out/'report_manifest.json').read_text(encoding='utf-8'))
    file=root/manifest['docx']; doc=Document(file)
    assert sha(file)==manifest['sha256']
    for name,digest in manifest['inputs'].items():assert sha(root/name)==digest,name
    assert len(doc.tables)==16 and len(doc.inline_shapes)==8
    numbers=[int(m.group(1)) for p in doc.paragraphs if p._element.findall('.//'+qn('m:oMath'))
             if (m:=re.search(r'\((\d+)\)$',p.text))]
    assert numbers==list(range(1,23)),numbers
    table_checks=[]
    for table,source in zip(doc.tables,expected['tables']):
        assert len(table.rows)==len(source['rows'])+1
        for row,raw in zip(table.rows[1:],source['rows']):
            # Count actual XML cells once; python-docx repeats merged-cell proxies.
            text=' '.join(''.join(c.itertext()) for c in row._tr.findall(qn('w:tc')))
            for value in raw:
                if isinstance(value,float):assert f'{value:,.2f}' in text,(source['number'],value)
        table_checks.append(source['number'])
    with zipfile.ZipFile(file) as z:
        embeds=Counter(hashlib.sha256(z.read(n)).hexdigest() for n in z.namelist() if n.startswith('word/media/'))
    assert embeds==Counter(sha(p) for p in (root/'figures/q12_revision').glob('*.png'))
    pdfs=list(args.render_dir.glob('*.pdf'));assert len(pdfs)==1
    pdf=fitz.open(pdfs[0]);assert len(pdf)==18
    first=pdf[0].get_text();assert '关键词' in first and '一 问题重述' not in first
    assert '问题重述' in pdf[1].get_text()
    page_checks=[]
    for i,page in enumerate(pdf):
        assert abs(page.rect.width-595.3)<1 and abs(page.rect.height-841.9)<1
        assert len(page.get_text().strip())>20
        for image in page.get_image_info():
            x0,y0,x1,y1=image['bbox'];assert x0>=70 and x1<=526 and y0>=68 and y1<=780
            assert re.search(r'图\s*\d+\s',page.get_text()),f'image without caption page {i+1}'
        page_checks.append({'page':i+1,'images':len(page.get_image_info())})
    sys.path.insert(0,str(Path.home()/'.codex/skills/math-modeling/tools/docx/scripts'))
    import paper_format as pf
    issues=pf.validate_paper_structure(doc,rendered_pages=len(pdf),min_content_units=0,
        min_equations=22,min_figures=8,min_tables=16,target_pages=0,official_max_pages=30)
    assert not issues,issues
    result={'pass':True,'docx_sha256':sha(file),'display_equations':22,'tables_values_checked':table_checks,
        'embedded_figures_equal_sources':True,'rendered_pages':page_checks,
        'paper_format_issues':issues,'scope_override':'User requested Q1/Q2 Word draft, not full four-question 20/25-page or 15000-word paper. No minimum length imposed.',
        'verifier_sha256':sha(__file__)}
    (out/'word_review.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
    print('PASS: Word hashes, 22 equation numbers, 16 numeric tables, 8 image hashes, 18 rendered pages and scope-adjusted structure checks.')


if __name__=='__main__':main()
