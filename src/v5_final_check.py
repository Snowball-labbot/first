"""Cross-artifact release checks; never substitute rendering for numeric audits."""
from pathlib import Path
import json,re,hashlib,zipfile
from docx import Document
from docx.oxml.ns import qn
from lxml import etree
from src.q12 import dump_json

OUT=Path('artifacts/v5')
def sha(path):
    p=Path(path);b=p.read_bytes()
    if p.suffix in {'.py','.md','.json','.tex'}:b=b.replace(b'\r\n',b'\n')
    return hashlib.sha256(b).hexdigest()

def main():
    manifests={name:json.loads((OUT/f'{name}.json').read_text(encoding='utf-8')) for name in
        ['verification','price_replay_audit','retained_reproduction','xlsx_audit','paper_manifest','render_audit','word_build']}
    for key in ['verification','price_replay_audit','retained_reproduction','xlsx_audit','render_audit']:assert manifests[key]['status']=='PASS'
    m=manifests['run_manifest'] if 'run_manifest' in manifests else json.loads((OUT/'run_manifest.json').read_text(encoding='utf-8'))
    for path,value in m['sources'].items():assert sha(path)==value,path
    for key in ['1','2','4-2']:assert (OUT/f'result{key}.xlsx').read_bytes()==Path(f'artifacts/v3/result{key}.xlsx').read_bytes()
    source=Path('reports/完整论文_V5.md');text=source.read_text(encoding='utf-8');doc=Document('reports/完整论文_V5.docx')
    p=manifests['paper_manifest'];assert len(doc.tables)==p['tables']==65 and len(doc.inline_shapes)==p['figures']==16
    headings=re.findall(r'^表 (\d+) ([^\n]+)\n\n\|',text,re.M);assert [int(x[0]) for x in headings]==list(range(1,66))
    paragraphs={p.text.strip() for p in doc.paragraphs}
    for n,title in headings:assert f'表 {n} {title}' in paragraphs,(n,title)
    table_numbers=set(re.findall(r'\d[\d,]*\.\d+','\n'.join(l for l in text.splitlines() if l.startswith('|'))))
    word_tables='\n'.join(cell.text for table in doc.tables for row in table.rows for cell in row.cells)
    for number in table_numbers:assert number in word_tables,number
    old=Document('reports/完整论文_V3.docx')._element.findall('.//'+qn('m:oMath'));new=doc._element.findall('.//'+qn('m:oMath'))
    def signature(x):return etree.tostring(x,method='c14n',exclusive=True)
    i=0
    for node in new:
        if i<len(old) and signature(node)==signature(old[i]):i+=1
    assert i==45 and len(new)==47,(i,len(new))
    def media(path):
        with zipfile.ZipFile(path) as z:return sorted(hashlib.sha256(z.read(n)).hexdigest() for n in z.namelist() if n.startswith('word/media/'))
    assert media('reports/完整论文_V4.docx')==media('reports/完整论文_V5.docx')
    render=manifests['render_audit'];assert sha('reports/完整论文_V5.docx')==render['docx_sha256'] and sha('reports/完整论文_V5.pdf')==render['pdf_sha256']
    assert render['body_and_references_pages_excluding_abstract']<=30 and render['abstract_pages']==1
    for ext in ['docx','pdf']:assert Path(f'reports/完整论文_V5.{ext}').stat().st_size<20*1024*1024
    for item in json.loads((OUT/'improvement.json').read_text(encoding='utf-8')):
        assert f'{item["v5_cost"]:,.2f}' in text and f'{item["saving"]:,.2f}' in text
        assert item['initial_soc_difference']==item['final_soc_difference']==0
    result=dict(status='PASS',tables=65,figures=16,old_native_math_preserved=45,new_native_equations=2,
        unique_numeric_table_values_checked=len(table_numbers),copied_workbooks_byte_identical=3,
        updated_workbook_cells_checked=sum(x['cells_compared'] for x in manifests['xlsx_audit']['updated']),
        pages=render['pages'],body_pages=render['body_and_references_pages_excluding_abstract'],
        figures_byte_identical_to_v4=True,
        sources={str(p):sha(p) for p in sorted(Path('src').glob('*v5*.py'))},
        outputs={str(p):sha(p) for p in sorted(Path('reports').glob('*V5*'))})
    dump_json(OUT/'final_check.json',result);print(json.dumps({k:v for k,v in result.items() if k not in ['sources','outputs']},ensure_ascii=False))
if __name__=='__main__':main()
