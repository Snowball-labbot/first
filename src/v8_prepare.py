"""V8 source binding, Q1 precision audit, and figure data profiles."""
from pathlib import Path
from decimal import Decimal, ROUND_HALF_UP
from dataclasses import replace
import hashlib
import json
import shutil
import openpyxl
import numpy as np
import pandas as pd
from scipy.io import loadmat, savemat
from src.q12 import read_inputs, optimize_day, audit_trace, B, interval_label
from src.deliver_q12 import blocks
from src.v3_model import solve
from src.q34_data import read_extended

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT/'artifacts/v8'
DATES = ['2025-03-20','2025-06-21','2025-09-23','2025-12-21']


def main():
    OUT.mkdir(exist_ok=True)
    data, source = read_inputs(ROOT.parent/'CUMCM2026Problems/C题')
    load, pv, price = [data[k] for k in ['q1_load','q1_pv','price']]
    frame, lp = solve(load, pv, price, 6000, equal=True)
    integer, mip = optimize_day(load, pv, price, 6000, (6000,'equal'))
    audit_trace(frame, initial=6000)
    cost = float(frame.plan_kwh@price)
    assert abs(cost-mip['cost']) < 1e-6
    assert abs(cost-35126.94858928963) < 1e-6
    experiments = {}
    for digits in [2,3,4]:
        rounded = np.round(frame.plan_kwh.to_numpy(), digits)
        experiments[f'purchase_rounded_{digits}dp_then_billed'] = float(rounded@price)
        _, meta = solve(np.round(load,digits), np.round(pv,digits), price,6000,equal=True)
        experiments[f'input_energy_rounded_{digits}dp_then_optimized'] = meta['objective']
    experiments['each_interval_bill_rounded_to_cents'] = float(sum(
        (Decimal(str(g))*Decimal(str(p))).quantize(Decimal('.01'),rounding=ROUND_HALF_UP)
        for g,p in zip(frame.plan_kwh,price)))
    _, rotated = solve(np.roll(load,1),np.roll(pv,1),np.roll(price,1),6000,equal=True)
    experiments['old_v7_rotated_midnight_boundary'] = rotated['objective']
    balance = (frame.plan_kwh + pv + frame.discharge_kwh-load-frame.charge_kwh-frame.spill_kwh)
    report = dict(status='PASS',cost_yuan=cost,rounded_total_yuan=round(cost,2),
        user_reference_yuan=35126.95,difference_yuan=35126.95-cost,
        linear_program=lp, mixed_integer_program=mip,
        max_balance_error_kwh=float(abs(balance).max()),
        initial_kwh=float(frame.soc_start_kwh.iloc[0]),final_kwh=float(frame.soc_end_kwh.iloc[-1]),
        rounding_and_alignment_experiments=experiments,
        source_hashes=source['input_files'],
        diagnosis='The V7 Q1-only circular shift moved the fixed midnight state boundary. Original row order gives 35126.94858928963. This is alignment, not rounding.',
        convention='All four questions treat an input timestamp as the right endpoint of its ten-minute interval. Template labels only control output row lookup; they do not reorder optimization inputs.')
    (OUT/'q1_precision_audit.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
    frame.insert(0,'slot',range(144));frame.insert(1,'interval',[interval_label(i) for i in range(144)])
    frame.to_csv(OUT/'q1_intervals.csv',index=False,float_format='%.17g')
    sensitivity=[]
    for eta in [.8,.85,.9,np.sqrt(.9),1.]:
        _, meta=optimize_day(load,pv,price,6000,(6000,'equal'),battery=replace(B,eta_c=eta,eta_d=eta))
        sensitivity.append(dict(single_efficiency=eta,roundtrip_efficiency=eta*eta,cost=meta['cost']))
    sensitivity=pd.DataFrame(sensitivity);sensitivity.to_csv(OUT/'q1_efficiency.csv',index=False)
    no_storage=float(price@np.maximum(load-pv,0))
    q1_summary=dict(cost=cost,no_storage_cost=no_storage,plan_kwh=float(frame.plan_kwh.sum()),
        charge_kwh=float(frame.charge_kwh.sum()),discharge_kwh=float(frame.discharge_kwh.sum()),
        savings_yuan=no_storage-cost,savings_percent=100*(1-cost/no_storage))
    (OUT/'q1_summary.json').write_text(json.dumps(q1_summary,indent=2),encoding='utf-8')
    d={k:v for k,v in loadmat(ROOT/'artifacts/v7/figure_data.mat').items() if not k.startswith('_')}
    old=loadmat(ROOT/'artifacts/v6/figure_data.mat')
    for key in ['validation_cost','validation_emergency','q2_cost','q3_validation']:
        d[key]=old[key]
    for key in ['load_kwh','pv_kwh','price','plan_kwh','charge_kwh','discharge_kwh','soc_start_kwh','soc_end_kwh']:
        d['q1_'+key]=frame[key].to_numpy()
    d['q1_efficiency']=sensitivity.roundtrip_efficiency.to_numpy()*100
    d['q1_sensitivity']=sensitivity.cost.to_numpy()/1e4
    d['q1_waterfall']=np.array([no_storage,sensitivity.cost.iloc[-1],cost])/1e4
    q2=pd.read_csv(ROOT/'artifacts/q12/ridge_q0.7_intervals.csv.gz') if (ROOT/'artifacts/q12/ridge_q0.7_intervals.csv.gz').exists() else None
    if q2 is None:
        candidates=list((ROOT/'artifacts').glob('*/ridge_q0.7_intervals.csv.gz'))
        assert len(candidates)==1,candidates
        q2=pd.read_csv(candidates[0])
    for key in ['load_kwh','pv_kwh','forecast_load_kwh','forecast_pv_kwh']:
        d['q2_'+key]=np.array([q2.loc[q2.date==date,key].to_numpy() for date in DATES])
    extended,_=read_extended(ROOT.parent/'CUMCM2026Problems/C题')
    d['q3_pv_actual']=extended['pv'][171]*6/1000
    d['q3_pv_issued']=extended['pv_issued'][171]*6/1000
    savemat(OUT/'figure_data.mat',d)
    profile={}
    for k,v in d.items():
        v=np.asarray(v);finite=v[np.isfinite(v)]
        profile[k]=dict(shape=list(v.shape),count=int(v.size),nonfinite=int(v.size-finite.size),
            minimum=float(finite.min()),maximum=float(finite.max()))
    (OUT/'figure_data_profile.json').write_text(json.dumps(profile,indent=2),encoding='utf-8')
    template=ROOT.parent/'CUMCM2026Problems/C题/附件/附件5/result1.xlsx'
    wb=openpyxl.load_workbook(template)
    assert not any(c.data_type=='f' for ws in wb for row in ws for c in row)
    for row in range(2,146):
        start=wb['计划购电量'].cell(row,1).value.split('-')[0].split('+')[0]
        hh,mm=map(int,start.split(':'));slot=(hh*6+mm//10)%144
        wb['计划购电量'].cell(row,2).value=float(frame.plan_kwh.iloc[slot])
    for i,b in enumerate(blocks(frame),2):
        for col,value in [(2,b[1]),(3,b[2])]:wb['充放电量'].cell(i,col).value=value
    wb['充放电量']['E2']=float(frame.soc_start_kwh.iloc[0])
    wb['充放电量']['E3']=float(frame.soc_end_kwh.iloc[-1])
    wb.save(OUT/'result1.xlsx');wb.close()
    for key in ['2','3','4-2','4-3']:
        shutil.copyfile(ROOT/f'artifacts/v7/result{key}.xlsx',OUT/f'result{key}.xlsx')
    print(json.dumps(report,ensure_ascii=False,indent=2))


if __name__=='__main__':
    main()
