"""Read-back verification of full Word and selected result workbooks."""
from pathlib import Path
import json,re,math,hashlib
from datetime import datetime
import openpyxl,fitz
from docx import Document
from docx.oxml.ns import qn

def sha(p):
    p=Path(p);b=p.read_bytes()
    if p.suffix in {'.py','.md','.tex','.json'}:b=b.replace(b'\r\n',b'\n')
    return hashlib.sha256(b).hexdigest()

def main():
    out=Path('artifacts/q34');payload=json.loads((out/'xlsx_payload.json').read_text(encoding='utf-8'));checks={}
    for key,p in payload.items():
        wb=openpyxl.load_workbook(out/f'result{key}.xlsx',data_only=False)
        maps={'计划购电量':'plan','充放电量':'storage','紧急购电量':'emergency','费用分解':'ledger'}
        if key!='4-2':maps['调整购电量']='adjusted'
        if p['versions']:maps['计划版本']='versions'
        count=0
        for sheet,field in maps.items():
            ws=wb[sheet]
            for i,row in enumerate(p[field],2):
                for j,expected in enumerate(row,1):
                    value=ws.cell(i,j).value
                    if j==1 and expected is not None:
                        assert isinstance(value,datetime) and value.strftime('%Y-%m-%d')==expected,(key,sheet,i,j,value)
                    elif isinstance(expected,(int,float)):
                        assert isinstance(value,(int,float)) and math.isclose(value,expected,rel_tol=1e-12,abs_tol=1e-7),(key,sheet,i,j,value,expected)
                    else:assert value==expected,(key,sheet,i,j,value,expected)
                    count+=1
        assert not any(c.data_type=='e' for ws in wb for row in ws for c in row)
        checks[key]={'cells_compared':count,'sha256':sha(out/f'result{key}.xlsx')}
    m=json.loads((out/'word_build.json').read_text(encoding='utf-8'));f=Path(m['docx']);d=Document(f)
    assert sha(f)==m['sha256']
    for f0,h in m['inputs'].items():assert sha(f0)==h,f0
    assert len(d.tables)==29 and len(d.inline_shapes)==16
    numbers=[int(x.group(1)) for p in d.paragraphs if p._element.findall('.//'+qn('m:oMath')) if (x:=re.search(r'\((\d+)\)$',p.text))]
    assert numbers==list(range(1,31))
    paragraphs=[p.text for p in d.paragraphs];code_files=['src/q12.py','src/q34_data.py','src/q34.py','src/q4_forecast.py','src/q4_run.py','src/verify_q12.py','src/verify_q34.py','src/export_q34_payload.py','scripts/export_xlsx.mjs','scripts/export_q34_xlsx.mjs']
    for name in code_files:
        start=paragraphs.index(name)+1;lines=Path(name).read_text(encoding='utf-8').splitlines()
        assert paragraphs[start:start+len(lines)]==[s or ' ' for s in lines],name
    pdf=fitz.open('.qa/full_paper/render/完整论文.pdf')
    appendix=next(i+1 for i,p in enumerate(pdf) if '附录 A 指定日期与结果文件' in p.get_text())
    assert appendix-1<=30
    assert all(abs(p.rect.width-595.28)<1 and abs(p.rect.height-841.89)<1 for p in pdf)
    assert f.stat().st_size<20_000_000
    result={'status':'PASS','workbooks':checks,'word':{'pages':len(pdf),'pages_before_appendices':appendix-1,'figures':16,'tables':29,'numbered_equations':30,'complete_source_files':len(code_files),'sha256':sha(f)}}
    (out/'delivery_verification.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8');print(json.dumps(result,ensure_ascii=False))

if __name__=='__main__':main()
