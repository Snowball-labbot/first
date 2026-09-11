"""LaTeX Markdown -> tested Pandoc OMML -> reference-style Word layout."""
from pathlib import Path
import hashlib
import json
import re
import sys
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT=Path(__file__).resolve().parents[1]
SKILL=Path.home()/'.codex/skills/math-modeling/tools/docx/scripts'
sys.path.insert(0,str(SKILL))
import paper_format as pf
from equations import markdown_to_docx, verify_conversion


def sha(path):
    path=Path(path);data=path.read_bytes()
    if path.suffix in {'.py','.md','.tex','.json'}:data=data.replace(b'\r\n',b'\n')
    return hashlib.sha256(data).hexdigest()


def font(style,size=12,bold=False,east='宋体'):
    style.font.name='Times New Roman';style.font.size=Pt(size);style.font.bold=bold
    style.font.color.rgb=RGBColor(0,0,0)
    style.font.italic=False
    rpr=style.element.get_or_add_rPr();rf=rpr.get_or_add_rFonts()
    for key,value in [('ascii','Times New Roman'),('hAnsi','Times New Roman'),('eastAsia',east)]:rf.set(qn('w:'+key),value)


def main():
    temp=ROOT/'.qa/full_paper';temp.mkdir(parents=True,exist_ok=True)
    source=ROOT/'reports/完整论文.md'
    reference=temp/'reference.docx'
    doc=pf.new_document(contest='cumcm')
    sec=doc.sections[0];sec.left_margin=sec.right_margin=Cm(2.6)
    sec.top_margin=sec.bottom_margin=Cm(2.54)
    for name in ['Normal','Body Text','First Paragraph']:
        if name in doc.styles:
            st=doc.styles[name];font(st);st.paragraph_format.line_spacing=1.15
            st.paragraph_format.space_after=Pt(5);st.paragraph_format.first_line_indent=Pt(24)
    doc.save(reference)
    intermediate=temp/'converted.docx'
    conversion=markdown_to_docx(source,intermediate,reference,overwrite=True)
    assert verify_conversion(intermediate)['passed']
    doc=Document(intermediate)
    sec=doc.sections[0];sec.page_width=Cm(21);sec.page_height=Cm(29.7)
    sec.left_margin=sec.right_margin=Cm(2.6);sec.top_margin=sec.bottom_margin=Cm(2.54)
    sec.footer_distance=Cm(1.2)
    for name,size,bold,east in [('Normal',11.5,False,'宋体'),('Body Text',11.5,False,'宋体'),
        ('First Paragraph',11.5,False,'宋体'),('Title',16,True,'黑体'),('Subtitle',10.5,False,'宋体'),
        ('Heading 1',14,True,'黑体'),('Heading 2',12,True,'黑体'),('Heading 3',11,True,'黑体'),
        ('Heading 4',11,True,'黑体'),('Caption',10.5,False,'宋体'),('Image Caption',10.5,False,'宋体')]:
        if name in doc.styles:font(doc.styles[name],size,bold,east)
    # Drop Pandoc's implicit image caption; explicit numbered Chinese caption follows.
    figure_titles=set(re.findall(r'!\[([^\]]+)\]\(',source.read_text(encoding='utf-8')))
    for p in list(doc.paragraphs):
        if p.text.strip() in figure_titles and not p._element.findall('.//'+qn('w:drawing')):
            p._element.getparent().remove(p._element)
    # Remove theme paragraph borders (including Word's default Title blue rule).
    for element in [doc._element,doc.styles.element]:
        for border in list(element.iter(qn('w:pBdr'))):border.getparent().remove(border)
    eqnum=0
    for i,p in enumerate(doc.paragraphs):
        t=p.text.strip();fmt=p.paragraph_format
        fmt.space_before=Pt(0);fmt.space_after=Pt(5);fmt.line_spacing=1.15
        fmt.first_line_indent=Pt(24);fmt.widow_control=True
        p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
        for run in p.runs:
            # Preserve semantic bold/italic but remove theme colors and font drift.
            run.font.color.rgb=RGBColor(0,0,0)
        if i==0:
            p.style=doc.styles['Title'];fmt.first_line_indent=Pt(0);fmt.space_after=Pt(10)
            p.alignment=WD_ALIGN_PARAGRAPH.CENTER;fmt.keep_with_next=True
        elif t=='第一问与第二问修订初稿':
            p.style=doc.styles['Subtitle'];p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            fmt.first_line_indent=Pt(0);fmt.keep_with_next=True;fmt.space_after=Pt(12)
        elif t=='摘 要':
            p.style=doc.styles['Heading 1'];p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            fmt.first_line_indent=Pt(0);fmt.keep_with_next=True;fmt.space_after=Pt(8)
        elif p.style.name.startswith('Heading'):
            # Markdown uses ## for main sections; map into proper Word outline levels.
            level=int(p.style.name.split()[-1])
            p.style=doc.styles['Heading '+str(max(1,level-1))]
            fmt.first_line_indent=Pt(0);fmt.space_before=Pt(10);fmt.space_after=Pt(7)
            fmt.keep_with_next=True
            p.alignment=WD_ALIGN_PARAGRAPH.CENTER if p.style.name=='Heading 1' else WD_ALIGN_PARAGRAPH.LEFT
            if t in ['一 问题重述','四 符号说明','参考文献'] or t.startswith('附录 '):fmt.page_break_before=True
        elif t.startswith('关键词：'):
            fmt.first_line_indent=Pt(0);fmt.space_before=Pt(6)
        elif re.match(r'^[图表]\s*\d+\s',t):
            p.style=doc.styles['Caption'];p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            fmt.first_line_indent=Pt(0);fmt.space_before=Pt(5);fmt.space_after=Pt(5)
            fmt.keep_with_next=t.startswith('表')
        elif p._element.findall('.//'+qn('w:drawing')):
            fmt.first_line_indent=Pt(0);fmt.space_after=Pt(0);fmt.keep_with_next=True
            p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        elif p._element.findall('.//'+qn('m:oMathPara')):
            eqnum+=1;fmt.first_line_indent=Pt(0);fmt.space_before=Pt(5);fmt.space_after=Pt(7)
            fmt.keep_together=True;p.alignment=WD_ALIGN_PARAGRAPH.LEFT
            # Convert display wrapper to a single inline math run between center/right tabs.
            mp=p._element.find(qn('m:oMathPara'))
            if mp is not None:
                om=mp.find(qn('m:oMath'))
                if om is not None:
                    index=p._element.index(mp);p._element.remove(mp)
                    lead=OxmlElement('w:r');lead.append(OxmlElement('w:tab'));p._element.insert(index,lead)
                    p._element.insert(index+1,om)
                    fmt.tab_stops.add_tab_stop(Cm(7.9),WD_TAB_ALIGNMENT.CENTER)
                    fmt.tab_stops.add_tab_stop(Cm(15.8),WD_TAB_ALIGNMENT.RIGHT)
                    run=p.add_run('\t('+str(eqnum)+')');pf.set_run_font(run,size=10.5)
                    for mr in om.findall('.//'+qn('m:r')):
                        rpr=mr.find(qn('w:rPr'))
                        if rpr is None:
                            rpr=OxmlElement('w:rPr')
                            mr.insert(1 if mr.find(qn('m:rPr')) is not None else 0,rpr)
                        sz=OxmlElement('w:sz');sz.set(qn('w:val'),'21');rpr.append(sz)
        elif t.startswith('[') or t.startswith('数学计算程序') or t.startswith('使用 Python'):
            fmt.first_line_indent=Pt(0)
            if t.startswith('['):p.alignment=WD_ALIGN_PARAGRAPH.LEFT
            for run in p.runs:run.font.size=Pt(10)
        if t.endswith('：'):fmt.keep_with_next=True
    for shape in doc.inline_shapes:
        factor=Cm(15.5)/shape.width;shape.width=Cm(15.5);shape.height=int(shape.height*factor)
    # Retain all original table numbers and values; set three-line borders and exact widths.
    for ti,table in enumerate(doc.tables,1):
        table.autofit=False;table.alignment=WD_TABLE_ALIGNMENT.CENTER
        pf._set_table_borders(table)
        columns=len(table.columns)
        if ti==1:widths=[4.0,4.2,7.6]
        elif ti==2:widths=[3.2,9.4,3.2]
        elif columns==6:widths=[2.9,2.25,2.75,2.9,2.25,2.75]
        elif ti==5:widths=[1.8,1.2,4.3,4.3,4.2]
        elif ti==6:widths=[2.6,3.3,3.3,3.3,3.3]
        elif ti==7:widths=[4.2,2.9,2.9,2.9,2.9]
        else:widths=[15.8/columns]*columns
        scale=15.8/sum(widths);widths=[v*scale for v in widths]
        for col,w in zip(table.columns,widths):col.width=Cm(w)
        for ri,row in enumerate(table.rows):
            trpr=row._tr.get_or_add_trPr();cant=OxmlElement('w:cantSplit');trpr.append(cant)
            if ri==0:trpr.append(OxmlElement('w:tblHeader'))
            for ci,cell in enumerate(row.cells):
                cell.width=Cm(widths[ci]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                for p in cell.paragraphs:
                    p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.space_before=Pt(3)
                    p.paragraph_format.space_after=Pt(3);p.paragraph_format.line_spacing=1.0
                    p.paragraph_format.keep_with_next=ri<len(table.rows)-1 and len(table.rows)<=10
                    p.alignment=WD_ALIGN_PARAGRAPH.CENTER
                    for run in p.runs:pf.set_run_font(run,size=8.5 if columns>=8 else 9.5,bold=ri==0)
                margins=OxmlElement('w:tcMar')
                for side in ['left','right']:
                    el=OxmlElement('w:'+side);el.set(qn('w:w'),'45');el.set(qn('w:type'),'dxa');margins.append(el)
                cell._tc.get_or_add_tcPr().append(margins)
        # The final row of six-column tables uses wide summary fields, as in the problem.
        if columns==6:
            row=table.rows[-1]
            if row.cells[0].text.startswith('00:00'):
                for left,right in [(1,2),(4,5)]:
                    text=row.cells[left].text;cell=row.cells[left].merge(row.cells[right]);cell.text=text
            elif row.cells[0].text=='全天购电量':
                amount=row.cells[1].text;cost=row.cells[3].text
                row.cells[0].merge(row.cells[1]).text='全天购电量  '+amount+' kWh'
                row.cells[2].merge(row.cells[5]).text='全天购电费  '+cost+' 元'
            for cell in row.cells:
                for p in cell.paragraphs:
                    p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.first_line_indent=Pt(0)
                    p.paragraph_format.space_before=Pt(3);p.paragraph_format.space_after=Pt(3)
                    for run in p.runs:pf.set_run_font(run,size=9.5)
        for cell in table.rows[0].cells:pf._set_cell_bottom(cell,'single','6')
    # Include complete computational sources, with natural Word line wrapping.
    code_files=['src/q12.py','src/q34_data.py','src/q34.py','src/q4_forecast.py','src/q4_run.py',
                'src/verify_q12.py','src/verify_q34.py','src/export_q34_payload.py',
                'scripts/export_xlsx.mjs','scripts/export_q34_xlsx.mjs']
    for filename in code_files:
        heading=doc.add_paragraph(filename,style='Heading 2')
        heading.paragraph_format.page_break_before=True
        for line in (ROOT/filename).read_text(encoding='utf-8').splitlines():
            cp=doc.add_paragraph();cp.paragraph_format.first_line_indent=Pt(0)
            cp.paragraph_format.space_before=Pt(0);cp.paragraph_format.space_after=Pt(0)
            cp.paragraph_format.line_spacing=1.0;cp.paragraph_format.widow_control=False
            run=cp.add_run(line or ' ');pf.set_run_font(run,size=7.5)
            run.font.name='Consolas'
    footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
    fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE')
    run=OxmlElement('w:r');text=OxmlElement('w:t');text.text='1';run.append(text);fld.append(run);footer._element.append(fld)
    doc.core_properties.title='预测误差与波动电价下的微网购电和储能调度'
    doc.core_properties.author='';doc.core_properties.last_modified_by=''
    # Low-level formatting helpers append properties; restore the OOXML schema order.
    sequences={
        'tblPr':'tblStyle tblpPr tblOverlap bidiVisual tblStyleRowBandSize tblStyleColBandSize tblW jc tblCellSpacing tblInd tblBorders shd tblLayout tblCellMar tblLook tblCaption tblDescription tblPrChange',
        'tcPr':'cnfStyle tcW gridSpan hMerge vMerge tcBorders shd noWrap tcMar textDirection tcFitText vAlign hideMark headers cellIns cellDel cellMerge tcPrChange',
    }
    for parent,sequence in sequences.items():
        order={qn('w:'+name):i for i,name in enumerate(sequence.split())}
        for node in doc._element.findall('.//'+qn('w:'+parent)):
            node[:]=sorted(node,key=lambda child:order.get(child.tag,999))
    final=ROOT/'reports/完整论文.docx';doc.save(final)
    expected=json.loads((ROOT/'artifacts/q34/paper_manifest.json').read_text(encoding='utf-8'))
    assert len(doc.tables)==expected['tables'] and len(doc.inline_shapes)==expected['figures'] and eqnum==expected['equations']
    files=[*(ROOT/'figures/full').glob('*.png'),source,ROOT/'reports/完整论文公式.tex',Path(__file__),*sorted((ROOT/'figures/q12_revision').glob('*.png')),
           ROOT/'artifacts/q12/run_manifest.json',ROOT/'src/q12.py']
    manifest={'docx':str(final.relative_to(ROOT)),'sha256':sha(final),'display_equations':eqnum,
        'hash_scheme':'SHA256; py/md/tex/json input line endings normalized to LF; binaries hashed verbatim',
        'native_math_objects':len(doc._element.findall('.//'+qn('m:oMath'))),'figures':len(doc.inline_shapes),'tables':len(doc.tables),
        'conversion_intermediate_verified':True,'pandoc_warnings':conversion['warnings'],
        'inputs':{str(f.relative_to(ROOT)):sha(f) for f in files},
        'layout':'A4; 2.54 cm vertical and 2.6 cm horizontal margins; editable OMML; three-line tables',
        'scope':'Four-question draft with full computational code appendix; team review required'}
    (ROOT/'artifacts/q34/word_build.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps({k:manifest[k] for k in ['docx','native_math_objects','figures','tables','pandoc_warnings']},ensure_ascii=False))


if __name__=='__main__':main()
