"""V12 editorial and pagination revision; preserve accepted figures and equations."""
from pathlib import Path
import json
import re
import shutil
from copy import deepcopy
from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph
from lxml import etree
from src.v11_document import set_font
from src.v9_document import maths, tables

ROOT=Path(__file__).resolve().parents[1]
PROGRAMS=['src/q12.py','src/q34_data.py','src/v3_model.py','src/v5_model.py',
    'src/v5b_value.py','src/v7_price.py','src/v7_dispatch.py','src/v12_reproduce.py',
    'src/v12_results.py','src/v10_prepare.py','scripts/v10_figures.m',
    'support/V12/references/roles/编程手/scripts/export_publication_figure.m',
    'support/V12/references/roles/编程手/scripts/audit_publication_figure.m']

# Source indices refer to the frozen V11 DOCX, not the evolving paragraph list.
EXPAND={
23:'计划层承担“提前买多少”的决策，执行层承担“当前放多少”的决策。两层共享同一储能状态，但优化时可用的信息不同：日前依赖预测，实时依据已揭示的供需。供需偏差通过紧急购电和弃电进入实际费用核算。',
66:'该等价性依赖于富余电量可自由弃用且不计额外弃电罚金。若加入必须消纳、售电收益或与充放电吞吐量相关的成本，应重新检查变换是否保持目标与约束。第一问另用含二元互斥变量的模型复核，最优费用与线性规划结果一致。',
71:'储电量接近上限时，继续低价购电的可利用空间受限；接近下限时，后续负载缺口需要外网补足。因此图中的充放电切换既受价格驱动，也受当前库存和未来光伏影响，不能仅以某一固定电价阈值解释全部调度行为。',
80:'无损储能与实际储能的费用差额包含重新优化后的计划变化，不能直接解释为损耗电量乘某一平均电价。该组敏感性结果说明，提高效率能够扩大跨时段转移的经济收益；具体收益仍取决于光伏富余、负载缺口和峰谷价格的共同分布。',
90:'上述特征兼顾日内周期、周内差异和近期水平变化。正则化用于约束相互相关的滞后特征，参数选择依据验证期误差完成；调度层的风险余量则依据验证费用选择，分别对应预测精度与运营成本两个目标。',
101:'余量进入名义负载后，会同时改变普通购电和储能预安排；执行阶段仍使用实际供需，不把余量当成必须消耗的实体负载。余量不足可能增加五倍价补购，余量过大则可能形成无法利用但仍需付费的合同电量。',
108:'不同历史日期对应不同的未来缺口路径，逐条递推后再取均值，可保留误差沿时间连续变化的特征。该处理给出有限历史样本下的库存价值近似；只有当前时段的动作实际执行，历史路径中的未来状态并不作为已知事实输入实时控制。',
122:'负载曲线表现出较稳定的日内水平与时段变化，滞后特征能够提供有效信息；光伏在部分日期仍出现峰值和局部波动偏差。图6用于观察这些时段差异，表6则同时统计全部日期，避免只依据少量展示日期判断全年预测效果。',
136:'紧急电量略增而紧急费用下降，说明优化目标与“最少补购电量”并不等价。普通合同已确定后，储能应优先替代边际成本较高的缺口，而非机械地优先满足时间上更早的缺口；费用和电量需要结合分析。',
140:'例如，6点更新时，0—6点已经执行的结果保持固定，以6点实际储电量为新的边界条件，对6—24点重新规划，只执行至12点。随后重复这一过程，使预测更新、合同修订与实际状态反馈形成闭环。',
149:'采用相对午夜合同的净结算后，同一未来区间即使经历多次修订，也仅在执行时确定最终保留、取消和新增量。这样可以使计划优化中的目标与实际账单一致，避免把尚未交付的中间计划变化重复计费。',
162:'因此，预报更新是否值得采用，应以新增信息经过调单规则后形成的净费用变化评价。本例中三次更新的组合在一月验证期确定，2—12月结果用于报告其后续表现；不同更新组合的全年排序不再用于回选策略。',
176:'至此，第三问延续第二问的“计划—执行”结构，同时改变了可用预报与未来补缺渠道。负载修正更新需求判断，滚动优化修改合同，库存价值控制分配有限储电量；三者依次衔接，并通过同一实际账单评价。',
196:'价格预测误差通过实际交易价格与计划时预期价格的差异影响费用，也会改变库存保留是否合算。历史残差范围用于展示这种误差，模型当前未在计划目标中显式加入价格场景或最坏情形惩罚，因而不将本方法称为具有价格风险保证的鲁棒优化。',
199:'低价区间的充电与高价区间的放电构成主要分布特征，但同一电价下仍可能存在不同功率方向。该差异与负载、光伏、既有合同和储能边界共同有关，分布图展示运行结果中的关联，不单独识别电价的因果效应。',
208:'两类增益区间对应各自固定策略下的配对日费用差，重采样以连续日期区块为单位，以保留短期相关性。该统计量用于衡量现有年度轨迹上的收益波动，不能替代跨年度、不同市场条件下的独立检验。',
220:'尤其在临近日出日落、价格突变或储能接近容量边界时，十分钟平均量可能掩盖区间内约束冲突。工程应用需要结合更高频数据检查功率可行性，并在新增约束后重新评估现有节费幅度。',
224:'推广时可保持能量平衡、状态递推和费用分项的基本结构，先依据当地交易规则确定可调整的合同范围，再利用独立历史时段校准预测与风险参数。新增成本或约束应进入优化和结算两端，并用相同边界条件比较实施前后的效果。',
}


def main():
    path=ROOT/'reports/完整论文_V11.docx'
    base=Document(path);d=Document(path);ps=list(d.paragraphs)
    md=path.with_suffix('.md').read_text(encoding='utf8')
    log=[]
    def replace_plain(index,new):
        nonlocal md
        old=ps[index].text
        assert md.count(old)==1,(index,old)
        ps[index].runs[0].text=new
        for r in ps[index].runs[1:]:r.text=''
        md=md.replace(old,new)
    replace_plain(2,ps[2].text+'模型参数在一月验证期确定，评价期内仅使用决策时已获得的信息，按十分钟区间连续跟踪库存并核算合同、调整和紧急补购费用。')
    replace_plain(35,'（1）附件中的00:10、00:20、…、24:00依次对应00:00—00:10、00:10—00:20、…、23:50—24:00区间，区间内功率按常值处理。四问均保持原始记录顺序，状态定义在区间边界，功率乘1/6小时换算为电量。结果工作簿同样按当天00:00—24:00的144个区间顺序填写。')
    replace_plain(232,'电量单位为kWh，费用单位为元，普通购电与紧急购电分列。滚动问题的指定时段和全天购电量指最终有效普通计划；午夜计划和最终调整计划填入题目指定工作表，各次更新记录保存在支撑材料的中间结果中。')
    for i,text in EXPAND.items():
        anchor=ps[i]
        element=OxmlElement('w:p')
        anchor._element.addnext(element)
        p=Paragraph(element,anchor._parent)
        p.style=d.styles['Body Text']
        p.add_run(text)
        old=anchor.text
        # Inline equations are represented differently in Markdown.
        if i in [66,108]:
            marker='因此可先用HiGHS' if i==66 else '可行集合A包含'
            lines=md.splitlines()
            pos=next(k for k,s in enumerate(lines) if marker in s)
            lines.insert(pos+1,'\n'+text)
            md='\n'.join(lines)
        else:
            assert md.count(old)==1,(i,old)
            md=md.replace(old,old+'\n\n'+text)
        log.append(dict(after_v11_paragraph=i,added=text))
    # Describe the existing density graphic without changing it.
    old=ps[202].text
    new=old+'电价与功率分箱宽度分别为0.05元/kWh和0.25 MW，色阶按活跃区间占比采用对数刻度。'
    ps[202].runs[0].text=new
    for r in ps[202].runs[1:]:r.text=''
    md=md.replace(old,new)
    # Put explanatory prose and short result tables ahead of large inline
    # figures. This lets text occupy the preceding page without shrinking plots.
    def move_before(elements, anchor):
        for element in elements:anchor.addprevious(element)
    move_before([ps[66]._element,ps[66]._element.getnext()],ps[64]._element)
    move_before([ps[71]._element,ps[71]._element.getnext()],ps[69]._element)
    move_before([ps[131]._element,ps[131]._element.getnext(),
                 ps[132]._element,ps[133]._element,ps[134]._element],ps[129]._element)
    move_before([ps[200]._element,ps[201]._element,ps[202]._element],ps[208]._element)
    # Mirror the same block order in the editable Markdown manuscript.
    def move_md(start, stop, anchor):
        nonlocal md
        a=md.index(start);b=md.index(stop,a)
        block=md[a:b];md=md[:a]+md[b:]
        pos=md.index(anchor);md=md[:pos]+block+md[pos:]
    move_md('为进行全年多次求解','### 5.2',
            next(s for s in md.splitlines() if s.startswith('![') and 'raw_q1_inputs' in s))
    move_md(ps[71].text,'表 3列出',
            next(s for s in md.splitlines() if s.startswith('![') and 'result_q1_dispatch' in s))
    move_md('表 7 第二问','### 7.4',
            next(s for s in md.splitlines() if s.startswith('![') and 'q2_cost' in s))
    image12=next(s for s in md.splitlines() if s.startswith('![') and 'price_dispatch' in s)
    move_md(image12,'### 9.3',ps[208].text)
    # Release hidden inherited chains; retain only local, meaningful pairs.
    appendix=False;code=False
    for p in d.paragraphs:
        t=p.text.strip();f=p.paragraph_format
        if t.startswith('附录 A'):appendix=True
        if t.startswith('附录 B'):code=True
        if code:continue
        is_drawing=bool(p._element.findall('.//'+qn('w:drawing')))
        is_math=bool(p._element.findall('.//'+qn('m:oMath')))
        heading=p.style.name.startswith('Heading')
        f.keep_with_next=False
        f.keep_together=False
        f.widow_control=True
        if heading:
            f.keep_with_next=True;f.keep_together=True
            f.space_before=Pt(8);f.space_after=Pt(4)
        elif is_drawing:
            f.keep_with_next=True;f.keep_together=True
            f.space_before=Pt(3);f.space_after=Pt(0);f.line_spacing=1
        elif p.style.name=='Caption':
            f.keep_with_next=t.startswith('表')
            f.keep_together=True;f.space_before=Pt(3);f.space_after=Pt(4)
        elif t.startswith('注：'):
            f.keep_together=False;f.space_before=Pt(0);f.space_after=Pt(3)
            f.line_spacing=1.05
        elif not is_math and not appendix and p.style.name!='Title':
            f.line_spacing=1.10;f.space_before=Pt(0);f.space_after=Pt(3)
            f.first_line_indent=Pt(24)
            for r in p.runs:
                set_font(r._element.get_or_add_rPr(),size=12)
            if t.startswith('[') or t.startswith('关键词'):
                f.first_line_indent=Pt(0)
        if is_math:
            f.keep_together=True
    # Formula lead-ins remain attached to the formula; formulas do not pull
    # the following figure or paragraph to the next page.
    before_code=True
    for p in d.paragraphs:
        if p.text.startswith('附录 B'):before_code=False
        if not before_code:break
        if p._element.findall('.//'+qn('m:oMath')) and re.search(r'\(\d+\)',p.text):
            prev=p._element.getprevious()
            if prev is not None and prev.tag==qn('w:p'):
                Paragraph(prev,p._parent).paragraph_format.keep_with_next=True
    for table in d.tables:
        for ri,row in enumerate(table.rows):
            for cell in row.cells:
                for p in cell.paragraphs:
                    p.paragraph_format.keep_with_next=(ri<len(table.rows)-1 and len(table.rows)<=9)
                    p.paragraph_format.keep_together=True
                    p.paragraph_format.space_after=Pt(0)
                    p.paragraph_format.space_before=Pt(0)
    for p in ps[2:9]:
        p.paragraph_format.line_spacing=1.40
        p.paragraph_format.space_after=Pt(6)
    for p in ps[226:231]:
        for r in p.runs:set_font(r._element.get_or_add_rPr(),size=11)
        p.paragraph_format.line_spacing=1.05
        p.paragraph_format.first_line_indent=Pt(0)
        p.paragraph_format.keep_together=True
    ps[231].paragraph_format.page_break_before=False
    # The electronic-format specification requires complete runnable source
    # listings and a supporting-file inventory in the appendix.
    start=ps[338]._element
    for element in list(d._element.body):
        if element is start:break
    tail=list(d._element.body)
    begin=tail.index(start)
    for element in tail[begin:]:
        if element.tag!=qn('w:sectPr'):d._element.body.remove(element)
    h=d.add_paragraph('附录 B 支撑文件与计算程序',style='Heading 1')
    h.paragraph_format.page_break_before=False
    inventory=[
        '支撑材料包含结果工作簿、逐时计算轨迹、计划更新中间结果、程序和绘图数据。题目原始附件不重复打包，重新计算时由使用者提供C题数据目录。',
        '（1）results/V12/附件5/：result1.xlsx、result2.xlsx、result3.xlsx、result4-2.xlsx、result4-3.xlsx，分别对应第一、二、三问和第四问的两个分支，保持原模板工作表结构。',
        '（2）artifacts/v8/q1_intervals.csv、artifacts/v5b/*selected_intervals.csv.gz：五个分支的完整逐时计算轨迹；artifacts/v5b/*versions.csv.gz保存第三问与第四问滚动分支的三次计划更新记录。',
        '（3）artifacts/v5/online_selection.csv保存历史月度选择结果；artifacts/v8/figure_data.mat与artifacts/v10/dispatch_distribution.mat保存本文图形的数值输入。',
        '（4）src/与scripts/内的完整程序见下文；MATLAB图形导出与检查函数置于support/V12/。运行依赖及命令见压缩包根目录README.md。',
        '复算入口为python -m src.v12_reproduce --data-root <C题目录>。省略--days时计算334天；--days 1只执行首日核验。该入口复算论文已选定策略，采用保存的历史月度选择结果，不重新开展参数搜索。Result填写入口为python -m src.v12_results --templates <附件5目录>，只生成原模板包含的工作表。',
        '绘图先运行python -m src.v10_prepare，再在MATLAB中加入scripts路径并调用v10_figures(pwd,fullfile(pwd,\'support\',\'V12\'))。以下列出全部计算和绘图源程序，代码中的路径均相对于工作目录。'
    ]
    md=md[:md.index('## 附录 B')]+ '## 附录 B 支撑文件与计算程序\n\n'+'\n\n'.join(inventory)+'\n\n'
    for text in inventory:
        p=d.add_paragraph(text,style='Body Text')
        p.paragraph_format.first_line_indent=Pt(21)
        p.paragraph_format.line_spacing=1.1;p.paragraph_format.space_after=Pt(4)
        for r in p.runs:set_font(r._element.get_or_add_rPr(),size=10.5)
    for filename in PROGRAMS:
        h=d.add_paragraph(filename,style='Heading 2')
        h.paragraph_format.keep_with_next=True
        text=(ROOT/filename).read_text(encoding='utf8').replace('\r\n','\n')
        language='matlab' if filename.endswith('.m') else 'python'
        md+='### '+filename+'\n\n'+f'\x60\x60\x60{language}\n'+text.rstrip()+'\n\x60\x60\x60\n\n'
        for line in text.splitlines():
            p=d.add_paragraph(style='Normal')
            p.paragraph_format.first_line_indent=Pt(0)
            p.paragraph_format.space_before=Pt(0);p.paragraph_format.space_after=Pt(0)
            p.paragraph_format.line_spacing=Pt(8.5)
            p.paragraph_format.keep_with_next=False;p.paragraph_format.keep_together=False
            p.paragraph_format.widow_control=False
            r=p.add_run(line)
            set_font(r._element.get_or_add_rPr(),western='Consolas',size=7.5)
    assert tables(d)==tables(base) and maths(d)==maths(base)
    assert [etree.tostring(x._inline,method='c14n') for x in d.inline_shapes]==[
        etree.tostring(x._inline,method='c14n') for x in base.inline_shapes]
    # Chinese hanging punctuation must not protrude into the required margin.
    successors=['w:topLinePunct','w:autoSpaceDE','w:autoSpaceDN','w:bidi',
        'w:adjustRightInd','w:snapToGrid','w:spacing','w:ind','w:contextualSpacing',
        'w:mirrorIndents','w:suppressOverlap','w:jc','w:textDirection',
        'w:textAlignment','w:textboxTightWrap','w:outlineLvl','w:divId',
        'w:cnfStyle','w:rPr','w:sectPr','w:pPrChange']
    for p in d.paragraphs:
        pr=p._element.get_or_add_pPr()
        tag=pr.find(qn('w:overflowPunct'))
        if tag is None:
            tag=OxmlElement('w:overflowPunct');pr.insert_element_before(tag,*successors)
        tag.set(qn('w:val'),'0')
    d.core_properties.author=''
    d.core_properties.last_modified_by=''
    dest=ROOT/'reports/完整论文_V12.docx'
    d.save(dest)
    dest.with_suffix('.md').write_text(re.sub(r'\n{3,}','\n\n',md).rstrip()+'\n',encoding='utf8')
    shutil.copyfile(ROOT/'reports/完整论文_V11公式.tex',ROOT/'reports/完整论文_V12公式.tex')
    out=ROOT/'artifacts/v12';out.mkdir(parents=True,exist_ok=True)
    (out/'editorial_changes.json').write_text(json.dumps(log,ensure_ascii=False,indent=2),encoding='utf8')
    print('V12',len(log),'substantive additions; figures/equations/tables preserved')


if __name__=='__main__':main()
