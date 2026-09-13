"""Independent content, font and pagination checks for the V11 editorial edition."""
from pathlib import Path
import json, re, hashlib, zipfile, collections
import fitz
from docx import Document
from docx.oxml.ns import qn
from lxml import etree
from src.v9_document import maths,tables
from src.v11_document import EDITS
ROOT=Path(__file__).resolve().parents[1]
def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
def media(p):
    with zipfile.ZipFile(p) as z:return {n:hashlib.sha256(z.read(n)).hexdigest() for n in z.namelist() if n.startswith('word/media/')}
def main():
    out=ROOT/'artifacts/v11';a=Document(ROOT/'reports/完整论文_V10.docx');b=Document(ROOT/'reports/完整论文_V11.docx')
    expected=[(i,EDITS.get(i,p.text)) for i,p in enumerate(a.paragraphs) if EDITS.get(i,p.text) is not None]
    moved=[item for item in expected if item[0] in [129,130]]
    expected=[item for item in expected if item[0] not in [129,130]]
    index=next(i for i,item in enumerate(expected) if item[0]==126)
    expected[index:index]=moved
    check={'documented_prose_edits_only':[x[1] for x in expected]==[p.text for p in b.paragraphs],
      '61_tables_content_identical':len(b.tables)==61 and tables(a)==tables(b),
      '52_native_math_objects_identical':len(maths(b))==52 and maths(a)==maths(b),
      'all_drawing_xml_identical':[etree.tostring(s._inline,method='c14n') for s in a.inline_shapes]==[etree.tostring(s._inline,method='c14n') for s in b.inline_shapes],
      'all_media_bytes_identical':media(ROOT/'reports/完整论文_V10.docx')==media(ROOT/'reports/完整论文_V11.docx'),
      'equation_tex_identical':(ROOT/'reports/完整论文_V10公式.tex').read_bytes()==(ROOT/'reports/完整论文_V11公式.tex').read_bytes(),
      'abstract_text_identical':[p.text for p in a.paragraphs[:9]]==[p.text for p in b.paragraphs[:9]],
      'references_text_identical':[p.text for p in a.paragraphs[231:236]]==[p.text for p in b.paragraphs if re.match(r'^\[[1-5]\] ',p.text)]}
    captions={k:[p.text for p in b.paragraphs if p.style.name=='Caption' and p.text.startswith(k)] for k in ['图','表']}
    check['sequential_unique_figure_captions']=[int(re.search(r'\d+',p)[0]) for p in captions['图']]==list(range(1,13))
    check['sequential_unique_table_captions']=[int(re.search(r'\d+',p)[0]) for p in captions['表']]==list(range(1,62))
    body='\n'.join(p.text for p in b.paragraphs if p.style.name!='Caption')
    check['all_objects_referenced']=all(re.search(fr'{k}\s*{n}(?!\d)',body) for k,maxn in [('图',12),('表',61)] for n in range(1,maxn+1))
    pre=b.paragraphs[:next(i for i,p in enumerate(b.paragraphs) if p.text=='十一 参考文献')]
    check['all_5_references_cited']=set(re.findall(r'\[([1-5])\]','\n'.join(p.text for p in pre)))==set('12345')
    pdf=fitz.open(ROOT/'reports/完整论文_V11.pdf');txt=[p.get_text() for p in pdf]
    appendix=next(i+1 for i,t in enumerate(txt) if '附录 A 四问' in t)
    fonts=collections.Counter()
    for page in pdf:
        for block in page.get_text('dict')['blocks']:
            for line in block.get('lines',[]):
                for s in line['spans']:fonts[(s['font'],round(s['size'],2))]+=len(s['text'])
    check['no_unintended_font_fallback']=all(f not in ['MS-Gothic','Calibri','Calibri-Bold','MicrosoftYaHei-Bold'] for f,size in fonts)
    check['abstract_one_page']='关键词' in txt[0] and '一 问题重述' in txt[1]
    check['body_and_references_within_30_pages']=appendix-2<=30
    check['no_blank_pages']=all(len(t.strip())>30 for t in txt)
    check['a4']=all(abs(p.rect.width-595.28)<1 and abs(p.rect.height-841.89)<1 for p in pdf)
    check['margins_at_least_2_5cm']=all(min(s.top_margin.cm,s.bottom_margin.cm,s.left_margin.cm,s.right_margin.cm)>=2.49 for s in b.sections)
    check['docx_pdf_under_20mb']=all((ROOT/f'reports/完整论文_V11.{ext}').stat().st_size<20*1024**2 for ext in ['docx','pdf'])
    pairs=json.loads((out/'figure_page_pairs.json').read_text(encoding='utf8'))
    check['all_12_figures_share_caption_page']=len(pairs)==12 and all(p['figure_page']==p['caption_page'] for p in pairs)
    check['figure12_note_same_page']=any('图12' in t.replace(' ','').replace('\n','') and '42,989' in t for t in txt)
    check['table6_note_same_page']=any('第二问全期预测误差' in t and '包含夜间零值' in t for t in txt)
    dims=[dict(figure=i+1,width_emu=s.width,height_emu=s.height) for i,s in enumerate(b.inline_shapes)]
    report=dict(passed=all(check.values()),checks=check,pages=len(pdf),body_and_references_pages=appendix-2,appendix_starts=appendix,paragraph_edits=len(EDITS),figure_dimensions=dims,
      fonts=[dict(font=f,size_pt=s,characters=n) for (f,s),n in fonts.most_common()],
      hashes={ext:sha(ROOT/f'reports/完整论文_V11.{ext}') for ext in ['docx','pdf','md']})
    (out/'verification.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf8')
    print(json.dumps({'passed':report['passed'],'checks':check,'pages':len(pdf),'body':appendix-2},ensure_ascii=False));assert report['passed']
if __name__=='__main__':main()
