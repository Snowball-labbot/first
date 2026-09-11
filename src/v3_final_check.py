"""Read back all populated delivery cells, without rewriting templates."""
from pathlib import Path
import json, math, hashlib
from datetime import datetime
import openpyxl
from openpyxl.utils.datetime import from_excel
from docx import Document

def main():
    root=Path('artifacts/v3')
    p=json.loads((root/'xlsx_payload.json').read_text(encoding='utf8'))
    count=0
    def check(actual, expected):
        nonlocal count
        count+=1
        if expected is None:
            assert actual is None, (actual,expected)
        elif isinstance(expected,str) and len(expected)==10 and expected[4]=='-':
            if isinstance(actual,(int,float)):actual=from_excel(actual)
            assert isinstance(actual,datetime) and actual.strftime('%Y-%m-%d')==expected,(actual,expected)
        elif isinstance(expected,(int,float)):
            assert isinstance(actual,(int,float)) and math.isclose(actual,expected,rel_tol=1e-11,abs_tol=1e-7),(actual,expected)
        else:assert actual==expected,(actual,expected)
    for key in ['3','4-2','4-3']:
        wb=openpyxl.load_workbook(root/f'result{key}.xlsx',read_only=True,data_only=True)
        old=openpyxl.load_workbook(Path('artifacts/v2')/f'result{key}.xlsx',read_only=True,data_only=True)
        for sheet,field in [('计划购电量','plan'),('调整购电量','adjusted'),('充放电量','storage'),('紧急购电量','emergency'),('费用分解','ledger'),('计划版本','versions')]:
            if key=='4-2' and field in ['adjusted','versions']:continue
            rows=p[key][field];ws=wb[sheet]
            assert tuple(next(ws.values))==tuple(next(old[sheet].values)),sheet
            for actual,expected in zip(ws.iter_rows(min_row=2,max_row=len(rows)+1,max_col=len(rows[0]),values_only=True),rows):
                for a,e in zip(actual,expected):check(a,e)
        wb.close();old.close()
    wb=openpyxl.load_workbook(root/'result1.xlsx',read_only=True,data_only=True)
    for actual,expected in zip(wb['计划购电量'].iter_rows(min_row=2,max_row=145,min_col=2,max_col=2,values_only=True),p['1']['plan']):check(actual[0],expected[0])
    for i,row in enumerate(p['1']['storage'],2):
        check(wb['充放电量'].cell(i,2).value,row[0]);check(wb['充放电量'].cell(i,3).value,row[1])
        if i<=3:check(wb['充放电量'].cell(i,5).value,row[2])
    wb.close()
    assert (root/'result2.xlsx').read_bytes()==Path('artifacts/v2/result2.xlsx').read_bytes()
    doc=Document('reports/完整论文_V3.docx')
    manifest=json.loads((root/'paper_manifest.json').read_text(encoding='utf8'))
    assert len(doc.tables)==manifest['tables'] and len(doc.inline_shapes)==manifest['figures']
    text=Path('reports/完整论文_V3.md').read_text(encoding='utf8')
    for item in json.loads((root/'specified_tables_manifest.json').read_text(encoding='utf8')):
        assert f'问题{item["question"]}（{item["date"]}）' in text
    result={'status':'PASS','cells_compared':count,'tables':len(doc.tables),'figures':len(doc.inline_shapes),'result2_byte_identical':True,'v2_workbook_headers_preserved':True,'word_sha256':hashlib.sha256(Path('reports/完整论文_V3.docx').read_bytes()).hexdigest()}
    (root/'final_check.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf8')
    print(json.dumps(result,ensure_ascii=False))

if __name__=='__main__':main()
