"""Read-only V7/V8 inventory and visual audit inputs before V9 editing."""
from pathlib import Path
import json,re,hashlib
import numpy as np
from scipy.io import loadmat
from docx import Document
from docx.oxml.ns import qn
from lxml import etree
from PIL import Image,ImageDraw
import fitz
ROOT=Path(__file__).resolve().parents[1];OUT=ROOT/'artifacts/v9';QA=ROOT/'.qa/v9'
def sha(p):return hashlib.sha256(Path(p).read_bytes()).hexdigest()
def sheets(paths,out,cols=4,thumb=(290,420)):
    for start in range(0,len(paths),cols*3):
        canvas=Image.new('RGB',(cols*(thumb[0]+14),3*(thumb[1]+30)),'#e8e8e8');d=ImageDraw.Draw(canvas)
        for j,path in enumerate(paths[start:start+cols*3]):
            im=Image.open(path).convert('RGB');im.thumbnail(thumb)
            x=(j%cols)*(thumb[0]+14)+7;y=(j//cols)*(thumb[1]+30)+24
            canvas.paste(im,(x,y));d.text((x,y-20),f'{start+j+1}: {path.stem}',fill='black')
        canvas.save(out/f'contact-{start//(cols*3)+1}.jpg')
def main():
    OUT.mkdir(exist_ok=True,parents=True);QA.mkdir(exist_ok=True,parents=True);result={}
    for ver in [7,8]:
        p=ROOT/f'reports/完整论文_V{ver}.docx';doc=Document(p);md=(ROOT/f'reports/完整论文_V{ver}.md').read_text(encoding='utf-8')
        figures=[];matches=re.findall(r'!\[([^\]]*)\]\(([^)]+)\)',md)
        for i,(shape,(title,path)) in enumerate(zip(doc.inline_shapes,matches),1):
            rid=shape._inline.find('.//'+qn('a:blip')).get(qn('r:embed'));blob=doc.part.related_parts[rid].blob
            figures.append(dict(number=i,title=title,path=path,width_cm=round(shape.width/360000,5),height_cm=round(shape.height/360000,5),media_sha256=hashlib.sha256(blob).hexdigest(),pixels=list(Image.open(ROOT/'reports'/path).size)))
        maths=[etree.tostring(n,method='c14n').decode() for n in doc._element.findall('.//'+qn('m:oMath'))]
        result[str(ver)]=dict(docx_sha256=sha(p),pdf_sha256=sha(p.with_suffix('.pdf')),markdown_sha256=sha(p.with_suffix('.md')),paragraphs=len(doc.paragraphs),heading_tree=[p.text for p in doc.paragraphs if p.style.name.startswith('Heading')],figures=figures,tables=[dict(number=i,rows=len(t.rows),cols=len(t.columns),headers=[c.text for c in t.rows[0].cells]) for i,t in enumerate(doc.tables,1)],math_count=len(maths),math_sha256=hashlib.sha256('\n'.join(maths).encode()).hexdigest())
    pdf=fitz.open(ROOT/'reports/完整论文_V8.pdf');paths=[];pages=[];out=QA/'baseline';out.mkdir(exist_ok=True)
    for i,p in enumerate(pdf):
        text=p.get_text();path=out/f'page-{i+1:02d}.png';p.get_pixmap(matrix=fitz.Matrix(1.25,1.25)).save(path);paths.append(path)
        blocks=p.get_text('dict')['blocks'];rects=[b['bbox'] for b in blocks if b['bbox'][1]>60 and b['bbox'][3]<800]
        pages.append(dict(page=i+1,characters=len(text.strip()),bottom=max([b[3] for b in rects],default=0),images=len(p.get_images()),head=text[:110],tail=text[-110:]))
    sheets(paths,out);(out/'full_text.txt').write_text('\n'.join(p.get_text() for p in pdf),encoding='utf-8')
    fout=QA/'baseline_figures';fout.mkdir(exist_ok=True)
    sheets([(ROOT/'reports'/x['path']).resolve() for x in result['8']['figures']],fout,cols=3,thumb=(490,345))
    data={k:v for k,v in loadmat(ROOT/'artifacts/v8/figure_data.mat').items() if not k.startswith('_')}
    profiles={k:dict(shape=list(v.shape),count=int(v.size),nonfinite=int((~np.isfinite(v)).sum()),minimum=float(np.nanmin(v)),maximum=float(np.nanmax(v))) for k,v in data.items()}
    result.update(base_commit='f4cdab0',reference_v7_commit='f8d78c1',pages=pages,data_profiles=profiles,figure_data_sha256=sha(ROOT/'artifacts/v8/figure_data.mat'),scope='V8 content frozen; only authorized figure 6 layout sentence may change')
    (OUT/'source_inventory.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps({v:{'figures':result[v]['figures'],'paragraphs':result[v]['paragraphs'],'math':result[v]['math_count'],'tables':len(result[v]['tables'])} for v in ['7','8']},ensure_ascii=False))
    print('V8 PDF pages:',len(pdf));print('Data profiles:',profiles)
if __name__=='__main__':main()
