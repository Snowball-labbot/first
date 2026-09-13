"""Fill official Attachment 5 copies from frozen, physically audited trajectories.

Run from the repository root:
python -m src.v12_results --templates PATH/附件5
No solver inputs, decisions, or original attachment files are modified.
"""
from pathlib import Path
from copy import copy
from datetime import datetime
import argparse
import hashlib
import json
import math
import numpy as np
import pandas as pd
import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from src.q12 import audit_trace, interval_label

ROOT = Path(__file__).resolve().parents[1]
SOURCES = {
    '1': 'artifacts/v8/q1_intervals.csv',
    '2': 'artifacts/v5b/q2_value_zero_selected_intervals.csv.gz',
    '3': 'artifacts/v5b/q3_value_release_selected_intervals.csv.gz',
    '4-2': 'artifacts/v5b/q42_value_selected_intervals.csv.gz',
    '4-3': 'artifacts/v5b/q43_value_release_selected_intervals.csv.gz',
}
VERSIONS = {
    '3': 'artifacts/v5b/q3_value_release_versions.csv.gz',
    '4-3': 'artifacts/v5b/q43_value_release_versions.csv.gz',
}
EXPECTED = {'1':35126.9485892896, '2':13991392.60398854,
            '3':13479283.31613424, '4-2':14765492.68144737,
            '4-3':14210881.010117685}


def events(g):
    values = g.emergency_kwh.to_numpy()
    result = []
    i = 0
    while i < 144:
        if values[i] <= 1e-7:
            i += 1
            continue
        j = i + 1
        while j < 144 and values[j] > 1e-7:
            j += 1
        result.append([interval_label(i).split('-')[0]+'-'+
                       interval_label(j-1).split('-')[1], float(values[i:j].sum())])
        i = j
    assert abs(sum(x[1] for x in result)-values.sum()) < 1e-5
    return result or [['无', 0.0]]


def trace(key):
    f = pd.read_csv(ROOT/SOURCES[key])
    if 'emergency_kwh' not in f:
        f['emergency_kwh'] = 0.
    if 'original_plan_kwh' not in f:
        f['original_plan_kwh'] = f.plan_kwh
    if key != '1':
        assert f.date.drop_duplicates().tolist() == list(
            pd.date_range('2025-02-01','2025-12-31').strftime('%Y-%m-%d'))
        assert len(f) == 48096
        assert np.array_equal(f.slot.to_numpy(), np.tile(np.arange(144),334))
    else:
        assert len(f) == 144 and np.array_equal(f.slot, np.arange(144))
    physical = audit_trace(f, initial=float(f.soc_start_kwh.iloc[0]))
    original = f.original_plan_kwh.to_numpy()
    final = f.plan_kwh.to_numpy()
    price = f.price.to_numpy()
    cancel = np.maximum(original-final,0)
    added = np.maximum(final-original,0)
    costs = np.column_stack([
        price*np.minimum(original,final), .5*price*cancel,
        1.5*price*added, 5*price*f.emergency_kwh.to_numpy()])
    total = costs.sum(axis=1)
    if 'total_cost_yuan' in f:
        assert np.max(np.abs(total-f.total_cost_yuan)) < 1e-7
    assert abs(total.sum()-EXPECTED[key]) < 1e-5
    for i, name in enumerate(['retained','cancelled','added','emergency']):
        f[name+'_bill'] = costs[:,i]
    f['actual_bill'] = total
    if key in ['1','2','4-2']:
        assert np.max(np.abs(final-original)) < 1e-7
    if key=='1':
        assert abs(f.soc_end_kwh.iloc[-1]-6000) < 1e-7
    return f, physical


def make_payload(key, f):
    times = [interval_label(i) for i in range(144)]
    if key == '1':
        storage = []
        for i,h in enumerate(range(0,24,4)):
            g = f.iloc[h*6:(h+4)*6]
            storage.append([f'{h}:00-{h+4}:00',float(g.charge_kwh.sum()),
                float(g.discharge_kwh.sum()), '0:00' if i==0 else '24:00' if i==1 else None,
                float(f.soc_start_kwh.iloc[0]) if i==0 else float(f.soc_end_kwh.iloc[-1]) if i==1 else None])
        payload = {
            '计划购电量': [['时间段','购电量'], *[[t,float(v)] for t,v in zip(times,f.plan_kwh)]],
            '充放电量': [['时间段','充电量','放电量','时刻','储电量'],*storage],
        }
    else:
        header = ['日期',*times,'全天购电量','全天购电费']
        payload = {'计划购电量':[header]}
        if key in VERSIONS:
            payload['调整购电量'] = [header]
        payload['充放电量'] = [['日期','时间段','充电量','放电量','时刻','储电量']]
        payload['紧急购电量'] = [['日期','购电时间段','购电量']]
        for date,g in f.groupby('date',sort=False):
            dt = datetime.strptime(date,'%Y-%m-%d')
            initial_fee = float((g.price*g.original_plan_kwh).sum())
            payload['计划购电量'].append([dt,*map(float,g.original_plan_kwh),
                float(g.original_plan_kwh.sum()),
                initial_fee if key in VERSIONS else float(g.actual_bill.sum())])
            if key in VERSIONS:
                payload['调整购电量'].append([dt,*map(float,g.plan_kwh),
                    float(g.plan_kwh.sum()),float(g.actual_bill.sum())])
            for i,h in enumerate(range(0,24,4)):
                block = g.iloc[h*6:(h+4)*6]
                payload['充放电量'].append([dt if i==0 else None,f'{h:02d}:00-{h+4:02d}:00',
                    float(block.charge_kwh.sum()),float(block.discharge_kwh.sum()),
                    '00:00' if i==0 else '24:00' if i==1 else None,
                    float(g.soc_start_kwh.iloc[0]) if i==0 else float(g.soc_end_kwh.iloc[-1]) if i==1 else None])
            payload['紧急购电量'].extend([[dt if i==0 else None,*e] for i,e in enumerate(events(g))])
        if key in VERSIONS:
            v = pd.read_csv(ROOT/VERSIONS[key])
            assert len(v)==334*(108+72+36)
            for date,g in f.groupby('date',sort=False):
                plan = g.original_plan_kwh.to_numpy().copy()
                for ver in [1,2,3]:
                    part = v[(v.date==date)&(v.version==ver)].sort_values('slot')
                    first = ver*36
                    assert np.array_equal(part.slot,np.arange(first,144))
                    assert np.max(np.abs(part.old_kwh-plan[first:])) < 1e-7
                    plan[first:] = part.new_kwh
                assert np.max(np.abs(plan-g.plan_kwh.to_numpy())) < 1e-7
    return payload


def write_book(template, dest, payload):
    wb = openpyxl.load_workbook(template)
    original_sheets = list(wb.sheetnames)
    assert not any(c.data_type=='f' for w in wb for row in w for c in row)
    assert not any(w.merged_cells.ranges for w in wb)
    # User chose midnight-to-midnight columns and expressly prohibited added
    # sheets or explanatory material inside Result. Keep only original sheets.
    payload={title:payload[title] for title in original_sheets}
    for title,rows in payload.items():
        if title not in ['计划购电量','调整购电量'] or len(rows[0])==2:
            rows[0]=[c.value for c in wb[title][1]]
        else:
            original_header=[c.value for c in wb[title][1]]
            rows[0][0]=original_header[0]
            rows[0][-2:]=original_header[-2:]
    for title, rows in payload.items():
        existed = title in wb
        ws = wb[title] if existed else wb.create_sheet(title)
        old_height = ws.max_row
        styles = [copy(ws.cell(2,j)._style) for j in range(1,len(rows[0])+1)]
        for row in ws:
            for cell in row:
                cell.value = None
        for i,row in enumerate(rows,1):
            for j,value in enumerate(row,1):
                cell = ws.cell(i,j)
                cell.value = value
                if i>old_height and existed:
                    cell._style = copy(styles[j-1])
                if isinstance(value,datetime):
                    cell.number_format = 'yyyy-mm-dd'
                elif isinstance(value,(float,int)):
                    cell.number_format = '0.00' if isinstance(value,float) else '0'
                if not existed:
                    cell.font = Font(name='宋体',size=11,bold=i==1,color='000000')
                    if i==1:
                        cell.fill = PatternFill('solid',fgColor='E6EDF1')
        if not existed:
            ws.column_dimensions['A'].width=24
            for j in range(2,len(rows[0])+1):
                ws.column_dimensions[openpyxl.utils.get_column_letter(j)].width=18
        if title=='口径说明':
            ws.column_dimensions['B'].width=96
            for row in ws:
                row[1].alignment=Alignment(wrap_text=True,vertical='center')
                ws.row_dimensions[row[0].row].height=36
    assert wb.sheetnames == original_sheets
    wb.properties.creator = 'Modeling'
    wb.properties.lastModifiedBy = 'Modeling'
    wb.properties.title = 'V12 附件5计算结果'
    wb.save(dest)
    wb.close()
    checked = openpyxl.load_workbook(dest,read_only=True,data_only=False)
    count=0
    for title, rows in payload.items():
        ws = checked[title]
        for actual,expected in zip(ws.iter_rows(min_row=1,max_row=len(rows),
                max_col=len(rows[0]),values_only=True),rows):
            for a,e in zip(actual,expected):
                count+=1
                if isinstance(e,(float,int)):
                    assert isinstance(a,(float,int)) and math.isclose(a,e,rel_tol=1e-12,abs_tol=1e-7),(title,a,e)
                else:
                    assert a==e,(title,a,e)
        assert all(c.value is None for row in ws.iter_rows(min_row=len(rows)+1) for c in row)
        assert not any(c.data_type in ['e','f'] for row in ws for c in row)
    checked.close()
    return dict(cells_compared=count,original_sheets=original_sheets,
        rows={k:len(v)-1 for k,v in payload.items()},
        sha256=hashlib.sha256(dest.read_bytes()).hexdigest())


def main():
    ap=argparse.ArgumentParser()
    ap.add_argument('--templates',type=Path,required=True)
    ap.add_argument('--output',type=Path,default=ROOT/'results/V12/附件5')
    args=ap.parse_args()
    args.output.mkdir(parents=True,exist_ok=True)
    report={'status':'PASS','workbooks':{}}
    for key in SOURCES:
        template=args.templates/f'result{key}.xlsx'
        before=hashlib.sha256(template.read_bytes()).hexdigest()
        f,physical=trace(key)
        payload=make_payload(key,f)
        item=write_book(template,args.output/template.name,payload)
        assert before==hashlib.sha256(template.read_bytes()).hexdigest()
        item.update(cost_yuan=float(f.actual_bill.sum()),physical=physical,
                    original_template_sha256=before,original_template_unchanged=True)
        report['workbooks'][key]=item
        print(key,item['cost_yuan'],item['cells_compared'],'cells PASS',flush=True)
    out=ROOT/'artifacts/v12'
    out.mkdir(parents=True,exist_ok=True)
    (out/'xlsx_audit.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf8')


if __name__=='__main__':
    main()
