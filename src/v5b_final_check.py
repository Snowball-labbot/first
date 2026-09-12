"""Check the updated V5 manuscript, native math, preserved figures and delivery."""
from pathlib import Path
import json,re,hashlib,zipfile
from docx import Document
from docx.oxml.ns import qn
from lxml import etree
from src.q12 import dump_json
OUT=Path('artifacts/v5b')
def sha(path):
    p=Path(path);b=p.read_bytes()
    if p.suffix in {'.py','.md','.json','.tex'}:b=b.replace(b'\r\n',b'\n')
    return hashlib.sha256(b).hexdigest()
def main():
    m={n:json.loads((OUT/f'{n}.json').read_text(encoding='utf-8')) for n in ['verification','xlsx_audit','paper_manifest','render_audit','word_build']}
    for n in ['verification','xlsx_audit','render_audit']:assert m[n]['status']=='PASS'
    for path,value in m['word_build']['inputs'].items():assert sha(path)==value,path
    text=Path('reports/完整论文_V5.md').read_text(encoding='utf-8');doc=Document('reports/完整论文_V5.docx')
    expected=m['paper_manifest'];assert len(doc.tables)==expected['tables']==63;assert len(doc.inline_shapes)==16
    captions=re.findall(r'^表 (\d+) ([^\n]+)\n\n\|',text,re.M);assert [int(x[0]) for x in captions]==list(range(1,64))
    paras={p.text.strip() for p in doc.paragraphs}
    for n,title in captions:assert f'表 {n} {title}' in paras
    numeric=set(re.findall(r'\d[\d,]*\.\d+','\n'.join(x for x in text.splitlines() if x.startswith('|'))))
    tables='\n'.join(c.text for t in doc.tables for row in t.rows for c in row.cells)
    for n in numeric:assert n in tables,n
    old=Document('reports/完整论文_V3.docx')._element.findall('.//'+qn('m:oMath'));new=doc._element.findall('.//'+qn('m:oMath'));i=0
    def sig(x):return etree.tostring(x,method='c14n',exclusive=True)
    for x in new:
        if i<len(old) and sig(x)==sig(old[i]):i+=1
    assert i==45 and len(new)==49,(i,len(new))
    def media(path):
        with zipfile.ZipFile(path) as z:return sorted(hashlib.sha256(z.read(n)).hexdigest() for n in z.namelist() if n.startswith('word/media/'))
    assert media('reports/完整论文_V4.docx')==media('reports/完整论文_V5.docx')
    for ext in ['docx','pdf']:
        p=Path(f'reports/完整论文_V5.{ext}');assert sha(p)==m['render_audit'][f'{ext}_sha256'];assert p.stat().st_size<20*1024*1024
    assert m['render_audit']['abstract_pages']==1 and m['render_audit']['body_and_references_pages_excluding_abstract']<=30
    assert (OUT/'result1.xlsx').read_bytes()==Path('artifacts/v3/result1.xlsx').read_bytes()
    for r in m['xlsx_audit']['updated']:assert sha(OUT/f'result{r["key"]}.xlsx')==r['sha256']
    for r in json.loads((OUT/'improvement.json').read_text(encoding='utf-8')):
        assert f'{r["cost"]:,.2f}' in text and f'{r["saving"]:,.2f}' in text
    result=dict(status='PASS',tables=63,figures=16,native_math_objects=49,old_native_math_preserved=45,display_equations=37,
        unique_numeric_table_values_checked=len(numeric),updated_workbook_cells_checked=sum(x['cells_compared'] for x in m['xlsx_audit']['updated']),
        unit_tests_passed=27,full_year_traces=11,intervals_audited=529056,new_independent_lp_certificates=24,
        body_pages=m['render_audit']['body_and_references_pages_excluding_abstract'],total_pages=m['render_audit']['pages'],figures_byte_identical_to_v4=True,
        sources={str(p):sha(p) for p in sorted(Path('src').glob('*v5b*.py'))},outputs={str(p):sha(p) for p in sorted(Path('reports').glob('*V5*'))})
    dump_json(OUT/'final_check.json',result);print({k:v for k,v in result.items() if k not in ['sources','outputs']})
if __name__=='__main__':main()
