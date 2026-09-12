"""Edit a copy of V8 OOXML; freeze text, table cells and native equations."""
from pathlib import Path
import json,re,hashlib,shutil,zipfile
from docx import Document
from docx.shared import Pt,Cm
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.part import Part
from docx.opc.packuri import PackURI
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from lxml import etree
ROOT=Path(__file__).resolve().parents[1]
OLD='图6按日期分列、按负载与光伏分行，同一行共用纵轴尺度。'
NEW='图6按日期分行、按负载与光伏分列，同一列共用纵轴尺度。'
def maths(d):return [etree.tostring(n,method='c14n') for n in d._element.iter(qn('m:oMath'))]
def tables(d):return [[[c.text for c in row.cells] for row in t.rows] for t in d.tables]
def main():
    source=ROOT/'reports/完整论文_V8.docx';d=Document(source);base=Document(source)
    md=(ROOT/'reports/完整论文_V8.md').read_text(encoding='utf8');assert md.count(OLD)==1
    manifest=json.loads((ROOT/'artifacts/v9/figure_manifest.json').read_text(encoding='utf8'))
    items={x['name']:x for x in manifest['figures']};matches=re.findall(r'!\[[^\]]*\]\(([^)]+)\)',md)
    replacement=0
    for p in d.paragraphs:
        if OLD in p.text:
            # Preserve run formatting and the rest of this paragraph.
            for r in p.runs:
                if OLD in r.text:r.text=r.text.replace(OLD,NEW);replacement+=1
    assert replacement==1
    for i,(shape,path) in enumerate(zip(d.inline_shapes,matches),1):
        name=Path(path).stem;item=items[name];shape.width=Cm(15.5);shape.height=Cm(item['height_cm'])
        blip=shape._inline.find('.//'+qn('a:blip'));rid=blip.get(qn('r:embed'))
        d.part.related_parts[rid]._blob=(ROOT/f'figures/v9/{name}.png').read_bytes()
        svg=Part(PackURI(f'/word/media/v9_{i}.svg'),'image/svg+xml',(ROOT/f'figures/v9/{name}_word.svg').read_bytes(),d.part.package)
        srid=d.part.relate_to(svg,RT.IMAGE)
        extlst=OxmlElement('a:extLst');ext=OxmlElement('a:ext');ext.set('uri','{96DAC541-7B7A-43D3-8B79-37D633B846F1}')
        sb=etree.Element('{http://schemas.microsoft.com/office/drawing/2016/SVG/main}svgBlip',nsmap={'asvg':'http://schemas.microsoft.com/office/drawing/2016/SVG/main'});sb.set(qn('r:embed'),srid);ext.append(sb);extlst.append(ext);blip.append(extlst)
    in_body=False;appendix=False;code_appendix=False
    for p in d.paragraphs:
        t=p.text.strip();fmt=p.paragraph_format
        if t=='一 问题重述':in_body=True
        if t.startswith('附录 '):appendix=True
        if t.startswith('附录 B'):code_appendix=True
        if not in_body:continue # Abstract typography is preserved exactly.
        if p.style.name.startswith('Heading'):
            fmt.keep_with_next=True;fmt.keep_together=True;fmt.space_before=Pt(9);fmt.space_after=Pt(5)
        elif p._element.findall('.//'+qn('w:drawing')):
            fmt.keep_with_next=True;fmt.keep_together=True;fmt.space_before=Pt(7);fmt.space_after=Pt(0);fmt.line_spacing=1
        elif re.match(r'^图\s*\d+',t):
            fmt.keep_with_next=False;fmt.keep_together=True;fmt.space_before=Pt(4);fmt.space_after=Pt(7)
        elif re.match(r'^表\s*\d+',t):fmt.keep_with_next=True;fmt.keep_together=True
        elif not appendix and not p._element.findall('.//'+qn('m:oMath')):
            fmt.line_spacing=1.12;fmt.space_after=Pt(3);fmt.widow_control=True
        elif code_appendix and p.style.name=='Normal':
            fmt.line_spacing=.98;fmt.space_after=Pt(0)
    for t in d.tables:
        pr=t._tbl.tblPr
        borders=pr.find(qn('w:tblBorders'))
        if borders is None:borders=OxmlElement('w:tblBorders');pr.append(borders)
        for child in list(borders):borders.remove(child)
        for edge in ['top','left','bottom','right','insideH','insideV']:
            n=OxmlElement('w:'+edge);n.set(qn('w:val'),'single' if edge in ['top','bottom'] else 'nil');n.set(qn('w:sz'),'8');n.set(qn('w:color'),'30353B');borders.append(n)
        for ri,row in enumerate(t.rows):
            rp=row._tr.get_or_add_trPr()
            if ri==0 and rp.find(qn('w:tblHeader')) is None:rp.append(OxmlElement('w:tblHeader'))
            if rp.find(qn('w:cantSplit')) is None:rp.append(OxmlElement('w:cantSplit'))
            for cell in row.cells:
                cp=cell._tc.get_or_add_tcPr();cb=cp.find(qn('w:tcBorders'))
                if cb is not None:cp.remove(cb)
                if ri==0:
                    cb=OxmlElement('w:tcBorders');n=OxmlElement('w:bottom');n.set(qn('w:val'),'single');n.set(qn('w:sz'),'4');n.set(qn('w:color'),'30353B');cb.append(n);cp.insert_element_before(cb,'w:shd','w:noWrap','w:tcMar','w:textDirection','w:tcFitText','w:vAlign','w:hideMark','w:headers','w:cellIns','w:cellDel','w:cellMerge','w:tcPrChange')
                for p in cell.paragraphs:
                    p.paragraph_format.space_before=Pt(2);p.paragraph_format.space_after=Pt(2)
                    if ri==0 or len(t.rows)<=8:
                        p.paragraph_format.keep_with_next=ri<len(t.rows)-1
    assert [p.text for p in d.paragraphs]==[p.text.replace(OLD,NEW) for p in base.paragraphs]
    assert tables(d)==tables(base);assert maths(d)==maths(base)
    dest=ROOT/'reports/完整论文_V9.docx';d.save(dest)
    md=md.replace(OLD,NEW).replace('../figures/v8/','../figures/v9/');(ROOT/'reports/完整论文_V9.md').write_text(md,encoding='utf8')
    for p in (ROOT/'reports').glob('*V8*.tex'):shutil.copyfile(p,p.with_name(p.name.replace('V8','V9')))
    audit={'base_commit':'f4cdab0','v7_layout_reference':'f8d78c1','paragraphs':len(d.paragraphs),'tables':len(d.tables),'native_math':len(maths(d)),'figures':len(d.inline_shapes),'text_unchanged_except_approved_sentence':True,'tables_identical':True,'native_math_identical':True,'svg_embedded':14,'color_exceptions':[3,6,8],'docx_sha256':hashlib.sha256(dest.read_bytes()).hexdigest()}
    (ROOT/'artifacts/v9/content_verification.json').write_text(json.dumps(audit,ensure_ascii=False,indent=2),encoding='utf8');print(audit)
if __name__=='__main__':main()
