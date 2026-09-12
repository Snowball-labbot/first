"""Artifact-level V8 checks: mathematical evidence, template cells and typesetting."""
from pathlib import Path
import hashlib
import json
import re
import fitz
import numpy as np
import pandas as pd
import openpyxl
from docx import Document
from docx.oxml.ns import qn
from PIL import Image
from src.q12 import audit_trace
from src.deliver_q12 import blocks

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'artifacts/v8'


def main():
    md=(ROOT/'reports/完整论文_V8.md').read_text(encoding='utf-8')
    doc=Document(ROOT/'reports/完整论文_V8.docx')
    pdf=fitz.open(ROOT/'reports/完整论文_V8.pdf')
    texts=[p.get_text() for p in pdf]
    assert not re.search(r'GRU|PyTorch|1\.002%|35,126\.85',md+'\n'.join(texts),re.I)
    assert not re.search(r'^### (3\.3|5\.[345])',md,re.M)
    assert not any(ord(c)<32 and c not in '\n\r\t' for c in md)
    assert len(doc.tables)==61 and len(doc.inline_shapes)==14
    assert len(doc._element.findall('.//'+qn('m:oMath')))==52
    captions=[int(v) for v in re.findall(r'!\[[^\]]*\]\([^)]+\)\n\n图\s*(\d+)',md)]
    assert captions==list(range(1,15)),captions
    table_numbers=[int(v) for v in re.findall(r'^表\s*(\d+)[^\n]+\n\n(?=\|)',md,re.M)]
    assert table_numbers==list(range(1,62)),table_numbers
    def table_text(number):
        return re.search(r'^表 '+str(number)+r' [^\n]+\n\n((?:\|[^\n]*\n?)+)',md,re.M).group(1).strip()
    assert table_text(3)==table_text(12)
    assert table_text(4)==table_text(13)
    appendix=next(i+1 for i,t in enumerate(texts) if '附录 A 四问' in t)
    code_start=next(i+1 for i,t in enumerate(texts) if '附录 B 支撑' in t)
    assert appendix-2<=30
    assert all(len(t.strip())>30 for t in texts)
    assert len(doc.sections[0].footer._element.findall('.//'+qn('w:fldSimple')))==1
    colors={s.get('color') for p in pdf for b in p.get_text('dict')['blocks'] if 'lines' in b
        for line in b['lines'] for s in line['spans'] if s['text'].strip()}
    assert colors=={0},colors
    figure_pages={}
    for number,title in re.findall(r'!\[[^\]]*\]\([^)]+\)\n\n图\s*(\d+)\s*([^\n]+)',md):
        exact=re.sub(r'\s+','','图'+number+title)
        found=[i for i,t in enumerate(texts) if exact in re.sub(r'\s+','',t)]
        assert len(found)==1,(number,found)
        page=pdf[found[0]];assert page.get_images(),number
        figure_pages[int(number)]=found[0]+1
    q=pd.read_csv(OUT/'q1_intervals.csv');physical=audit_trace(q,initial=6000.)
    q1cost=float(q.price@q.plan_kwh)
    assert round(q1cost,2)==35126.95
    assert abs(q.soc_end_kwh.iloc[-1]-6000)<1e-7
    audit=json.loads((OUT/'q1_precision_audit.json').read_text(encoding='utf-8'))
    assert abs(audit['mixed_integer_program']['dual_bound']-q1cost)<1e-6
    source=ROOT.parent/'CUMCM2026Problems/C题/附件/附件5/result1.xlsx'
    wb=openpyxl.load_workbook(OUT/'result1.xlsx',data_only=False)
    template=openpyxl.load_workbook(source,data_only=False)
    assert wb.sheetnames==template.sheetnames
    chronological=np.empty(144);seen=[]
    for row in range(2,146):
        label=wb['计划购电量'].cell(row,1).value
        assert label==template['计划购电量'].cell(row,1).value
        start=label.split('-')[0].split('+')[0];hh,mm=map(int,start.split(':'));slot=(hh*6+mm//10)%144
        seen.append(slot);actual=wb['计划购电量'].cell(row,2).value;chronological[slot]=actual
        assert abs(actual-q.plan_kwh.iloc[slot])<1e-7
    assert sorted(seen)==list(range(144))
    assert abs(float(q.price@chronological)-q1cost)<1e-7
    for i,b in enumerate(blocks(q),2):
        for col,value in [(2,b[1]),(3,b[2])]:
            assert abs(wb['充放电量'].cell(i,col).value-value)<1e-7
    assert wb['充放电量']['E2'].value==6000 and wb['充放电量']['E3'].value==6000
    assert not any(c.data_type=='f' for ws in wb for row in ws for c in row)
    wb.close();template.close()
    for key in ['2','3','4-2','4-3']:
        assert (OUT/f'result{key}.xlsx').read_bytes()==(ROOT/f'artifacts/v7/result{key}.xlsx').read_bytes()
    image_quality={}
    for image_path in re.findall(r'!\[[^\]]*\]\(([^)]+)\)',md):
        path=(ROOT/'reports'/image_path).resolve()
        for suffix in ['.png','.svg','.pdf','.fig']:
            assert path.with_suffix(suffix).exists()
        assert (path.parent/'_qa'/f'{path.stem}_grayscale.png').exists()
        with Image.open(path) as im:
            image_quality[path.stem]=dict(pixels=list(im.size),effective_dpi_at_15_5cm=im.width/(15.5/2.54))
            assert im.width/(15.5/2.54)>=500
        check=json.loads(path.with_name(path.stem+'_audit.json').read_text(encoding='utf-8'))
        assert check['all_text_black'] and not check['design_issues'],(path.name,check)
    manifest=json.loads((OUT/'word_build.json').read_text(encoding='utf-8'))
    for name,expected in manifest['inputs'].items():
        raw=(ROOT/name).read_bytes()
        if Path(name).suffix in {'.py','.md','.tex','.json'}:raw=raw.replace(b'\r\n',b'\n')
        assert hashlib.sha256(raw).hexdigest()==expected,name
    result=dict(status='PASS',pdf_pages=len(pdf),body_and_references_excluding_abstract=appendix-2,
        appendix_starts=appendix,program_appendix_starts=code_start,figures=14,tables=61,native_math=52,
        numbered_equations=27,q1_cost_yuan=q1cost,q1_physical=physical,
        q1_lp_milp_agree=True,q1_workbook_all_144_cells_checked=True,
        q1_body_and_appendix_tables_equal=True,other_four_workbooks_byte_identical=True,
        single_footer_page_field=True,all_pdf_text_black=True,removed_content_absent=True,
        figure_pages=figure_pages,image_quality=image_quality,
        hashes={f'reports/完整论文_V8{s}':hashlib.sha256((ROOT/f'reports/完整论文_V8{s}').read_bytes()).hexdigest()
            for s in ['.md','.docx','.pdf']},
        visual_review='See visual_review.json for final rendered-page inspection; deterministic checks do not replace it.')
    (OUT/'verification.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps({k:result[k] for k in ['status','pdf_pages','body_and_references_excluding_abstract','appendix_starts','q1_cost_yuan','all_pdf_text_black']},ensure_ascii=False))


if __name__=='__main__':main()
