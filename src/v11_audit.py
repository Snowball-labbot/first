"""Read V10 and user supplied reference PDFs without changing their contents."""
from pathlib import Path
import json,hashlib,collections,argparse
import fitz
from docx import Document
from docx.oxml.ns import qn
from src.v9_inventory import sheets
ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'artifacts/v11';QA=ROOT/'.qa/v11';OUT.mkdir(parents=True,exist_ok=True);QA.mkdir(parents=True,exist_ok=True)
def inspect_pdf(path,name):
 p=fitz.open(path);dest=QA/name;dest.mkdir(exist_ok=True);fonts=collections.Counter();paths=[]
 for i,page in enumerate(p):
  for b in page.get_text('dict')['blocks']:
   for l in b.get('lines',[]):
    for s in l['spans']:fonts[(s['font'],round(s['size'],2))]+=len(s['text'])
  img=dest/f'page-{i+1:02}.png';page.get_pixmap(matrix=fitz.Matrix(1,1)).save(img);paths.append(img)
 (dest/'full_text.txt').write_text('\n'.join(f'\n=== PAGE {i+1} ===\n'+page.get_text() for i,page in enumerate(p)),encoding='utf8');sheets(paths,dest)
 return dict(name=Path(path).name,sha256=hashlib.sha256(Path(path).read_bytes()).hexdigest(),pages=len(p),fonts=[dict(font=k[0],size=k[1],characters=v) for k,v in fonts.most_common(20)])
def main():
 ap=argparse.ArgumentParser();ap.add_argument('--references',nargs=2);args=ap.parse_args()
 d=Document(ROOT/'reports/完整论文_V10.docx');paras=[];fonts=collections.Counter()
 for i,p in enumerate(d.paragraphs):
  paras.append(dict(index=i,style=p.style.name,text=p.text,math=bool(p._element.findall('.//'+qn('m:oMath'))),drawing=bool(p._element.findall('.//'+qn('w:drawing')))))
  for r in p.runs:
   rf=r._element.find('.//'+qn('w:rFonts'));fonts[(r.font.name,str(r.font.size.pt if r.font.size else None),rf.get(qn('w:eastAsia')) if rf is not None else None)]+=len(r.text)
 inv=dict(source_commit='054e2e7',docx_sha256=hashlib.sha256((ROOT/'reports/完整论文_V10.docx').read_bytes()).hexdigest(),paragraphs=paras,tables=len(d.tables),math=len(d._element.findall('.//'+qn('m:oMath'))),figures=[dict(index=i+1,width=s.width,height=s.height) for i,s in enumerate(d.inline_shapes)],direct_fonts=[dict(font=k[0],size=k[1],eastAsia=k[2],characters=v) for k,v in fonts.most_common()])
 (QA/'paragraphs.txt').write_text('\n'.join(f"[{x['index']}] {x['style']} {'MATH' if x['math'] else ''} {x['text']}" for x in paras),encoding='utf8')
 inv['pdf']=inspect_pdf(ROOT/'reports/完整论文_V10.pdf','v10')
 inv['references']=[inspect_pdf(p,n) for p,n in zip(args.references,['C050','E010'])]
 (OUT/'source_audit.json').write_text(json.dumps(inv,ensure_ascii=False,indent=2),encoding='utf8')
 print(json.dumps({k:v for k,v in inv.items() if k!='paragraphs'},ensure_ascii=False))
if __name__=='__main__':main()
