"""Independent delivery checks against frozen V8 and saved figure contracts."""
from pathlib import Path
import json,hashlib,re,zipfile
import numpy as np
from scipy.io import loadmat
from docx import Document
from lxml import etree
import fitz
from src.v9_document import OLD,NEW,maths,tables
ROOT=Path(__file__).resolve().parents[1]
def main():
 a=Document(ROOT/'reports/完整论文_V8.docx');b=Document(ROOT/'reports/完整论文_V9.docx')
 checks={}
 checks['paragraphs_equal_with_one_approved_layout_sentence']=[p.text.replace(OLD,NEW) for p in a.paragraphs]==[p.text for p in b.paragraphs]
 checks['all_61_tables_equal']=tables(a)==tables(b)
 checks['all_52_native_math_objects_equal']=maths(a)==maths(b)
 ma=(ROOT/'reports/完整论文_V8.md').read_text(encoding='utf8');mb=(ROOT/'reports/完整论文_V9.md').read_text(encoding='utf8')
 checks['markdown_content_equal']=ma.replace(OLD,NEW).replace('../figures/v8/','../figures/v9/')==mb
 inv=json.loads((ROOT/'artifacts/v9/source_inventory.json').read_text(encoding='utf8'));mf=json.loads((ROOT/'artifacts/v9/figure_manifest.json').read_text(encoding='utf8'));d=loadmat(ROOT/'artifacts/v8/figure_data.mat',squeeze_me=True)
 checks['figure_source_file_unchanged']=mf['source_sha256']==inv['figure_data_sha256']
 checks['all_figure_source_arrays_equal']=all(hashlib.sha256(np.asarray(d[k]).tobytes()).hexdigest()==v for f in mf['figures'] for k,v in f['source_hashes'].items())
 checks['restored_three_v7_proportions']=all(abs(b.inline_shapes[i-1].height/360000-inv['7']['figures'][i-1]['height_cm'])<1e-4 for i in [3,6,8])
 checks['svg_for_all_figures']=len(b.part._element.findall('.//{http://schemas.microsoft.com/office/drawing/2016/SVG/main}svgBlip'))==14
 with zipfile.ZipFile(ROOT/'reports/完整论文_V9.docx') as z:
  for name in z.namelist():
   if name.endswith(('.xml','.rels','.svg')):etree.fromstring(z.read(name))
 checks['ooxml_and_svg_well_formed']=True
 p=fitz.open(ROOT/'reports/完整论文_V9.pdf');q=fitz.open(ROOT/'reports/完整论文_V8.pdf')
 wa=p[0].get_text('words');wb=q[0].get_text('words')
 checks['abstract_text_and_layout_preserved']=p[0].get_text()==q[0].get_text() and len(wa)==len(wb) and max(abs(x-y) for r,s in zip(wa,wb) for x,y in zip(r[:4],s[:4]))<.2
 checks['single_column_a4']=all(abs(x.rect.width-595.3)<1 and abs(x.rect.height-841.9)<1 for x in p)
 checks['docx_pdf_under_20mb']=all((ROOT/f'reports/完整论文_V9.{ext}').stat().st_size<20*1024*1024 for ext in ['docx','pdf'])
 captions={}
 for i,page in enumerate(p):
  for num in re.findall(r'图\s*(\d+)\s',page.get_text()):
   if 1<=int(num)<=14:captions.setdefault(num,[]).append(i+1)
 report=dict(checks=checks,passed=all(checks.values()),pages=len(p),counts=dict(paragraphs=len(b.paragraphs),tables=len(b.tables),math=len(maths(b)),figures=len(b.inline_shapes)),caption_pages=captions,notes=['图3、6、8配色按用户指令保留V7，是统一色表的指定例外。','全年52,560点价格热图作为高密度栅格嵌入SVG/PDF；坐标、文字及其他图元为矢量。','本次只作视觉重构，未重新拟合或优化数学模型。'])
 (ROOT/'artifacts/v9/verification.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf8');print(json.dumps(report,ensure_ascii=False));assert report['passed']
if __name__=='__main__':main()
