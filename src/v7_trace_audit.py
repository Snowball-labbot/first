"""Read raw workbooks independently and audit the four retained strategies."""
from pathlib import Path
import argparse,json,hashlib
import numpy as np
import pandas as pd
import openpyxl
from src.q12 import audit_trace,dump_json
from src.v7_delivery import SOURCES
OUT=Path('artifacts/v7')
def main():
    ap=argparse.ArgumentParser();ap.add_argument('--data-root',required=True);a=ap.parse_args();root=Path(a.data_root,'附件')
    books=[];raw={}
    for i in range(1,5):
        p=root/f'附件{i}.xlsx';w=openpyxl.load_workbook(p,read_only=False,data_only=False)
        books.append(dict(file=p.name,sha256=hashlib.sha256(p.read_bytes()).hexdigest(),sheets=[]))
        for ws in w:
            rows=list(ws.values);assert not any(c.data_type=='f' for row in ws for c in row)
            books[-1]['sheets'].append(dict(name=ws.title,rows=ws.max_row,columns=ws.max_column,hidden_rows=sum(bool(r.hidden) for r in ws.row_dimensions.values()),hidden_columns=sum(bool(r.hidden) for r in ws.column_dimensions.values()),state=ws.sheet_state))
            if i==1:raw['price']=np.array([r[1] for r in rows[1:]],float);raw['q1load']=np.array([r[2] for r in rows[1:]],float)/6;raw['q1pv']=np.array([r[3] for r in rows[1:]],float)/6
            if i==2:raw['load' if ws==w.worksheets[0] else 'pv']=np.array([r[1:] for r in rows[1:]],float)/6
            if i==4:raw['actual_price']=np.array([r[1:] for r in rows[1:]],float)
        w.close()
    records=[]
    for key,path in SOURCES.items():
        f=pd.read_csv(path);physical=audit_trace(f,initial=6000 if key=='1' else 9902.811287870369)
        p=raw['price'] if key=='1' else np.tile(raw['price'],334) if key in ['2','3'] else raw['actual_price'][31:].ravel()
        np.testing.assert_allclose(f.price,p,atol=1e-12,rtol=0)
        for column,name in [('load_kwh','load'),('pv_kwh','pv')]:np.testing.assert_allclose(f[column],raw['q1'+name] if key=='1' else raw[name][31:].ravel(),atol=1e-10,rtol=0)
        if key=='1':fee=p*f.plan_kwh
        else:
            assert len(f)==48100-4 # 334*144
            assert list(f.date.unique())==list(pd.date_range('2025-02-01','2025-12-31').strftime('%Y-%m-%d'))
            np.testing.assert_array_equal(f.slot,np.tile(np.arange(144),334))
            original=f.original_plan_kwh;final=f.plan_kwh
            fee=p*(np.minimum(original,final)+.5*np.maximum(original-final,0)+1.5*np.maximum(final-original,0)+5*f.emergency_kwh)
            np.testing.assert_allclose(fee,f.total_cost_yuan,atol=1e-7,rtol=0)
        records.append(dict(question=key,cost=float(fee.sum()),intervals=len(f),physical=physical,trace_sha256=hashlib.sha256(Path(path).read_bytes()).hexdigest()))
    dump_json(OUT/'trace_audit.json',dict(status='PASS',raw_workbooks=books,strategies=records,intervals=sum(r['intervals'] for r in records)))
    print('Raw-input and bill audit passed',[(r['question'],r['cost']) for r in records])
if __name__=='__main__':main()
