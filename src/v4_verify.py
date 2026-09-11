"""Verify editorial preservation, evidence, actual Word objects and rendered pages."""
from pathlib import Path
import hashlib
import json
import re
import zipfile
import pandas as pd
import pymupdf as fitz
from docx import Document
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[1]

def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
def load(p):return json.loads((ROOT/p).read_text(encoding='utf-8'))

def main():
    out=ROOT/'artifacts/v4';old=Document(ROOT/'reports/完整论文_V3.docx');new=Document(ROOT/'reports/完整论文_V4.docx')
    a=(ROOT/'reports/完整论文_V3.md').read_text(encoding='utf-8');b=(ROOT/'reports/完整论文_V4.md').read_text(encoding='utf-8')
    def tables(doc):return [[[c.text for c in r.cells] for r in t.rows] for t in doc.tables]
    assert tables(old)==tables(new),'Changed Word table cells'
    assert re.findall(r'\$\$(.*?)\$\$',a,re.S)==re.findall(r'\$\$(.*?)\$\$',b,re.S)
    def math(doc):return [''.join(e.itertext()) for e in doc._element.findall('.//'+qn('m:oMath'))]
    assert math(old)==math(new),'Changed native mathematical objects'
    assert len(new.tables)==62 and len(new.inline_shapes)==16 and len(math(new))==45
    def media(doc):
        return [hashlib.sha256(doc.part.related_parts[s._inline.graphic.graphicData.pic.blipFill.blip.embed].blob).hexdigest() for s in doc.inline_shapes]
    assert media(old)==media(new)[1:],'Changed inherited figures'
    assert media(new)[0]==sha(ROOT/'figures/v4/modeling_overview.png')
    pre=[]
    for p in new.paragraphs:
        if p.text.startswith('附录 A'):break
        pre.append(p)
    main_text='\n'.join(p.text for p in pre)
    for marker in ['**','`','$$','[待补充]']:
        assert marker not in main_text,marker
    assert not re.search(r'^\s*#{1,6}\s',main_text,re.M)
    for h in ['1.1 研究背景','1.2 问题的提出',*[f'2.{i} 问题{c}的分析' for i,c in enumerate('一二三四',1)]]:
        assert h in main_text,h
    abstract=main_text.split('摘 要',1)[1].split('一 问题重述',1)[0]
    methods=['确定性多期线性规划模型','岭回归预测与分位数风险校准模型','非对称结算下的滚动时域线性规划模型','残差GRU集成电价预测模型','真实电价名义最优下界']
    bold=''.join(r.text for p in pre[:10] for r in p.runs if r.bold)
    for term in methods:assert term in bold,term
    captions=[p.text for p in pre if p.style.name=='Caption']
    nums=[int(re.match(r'图\s*(\d+)',t)[1]) for t in captions if re.match(r'图\s*(\d+)',t)]
    assert nums==list(range(1,17)),nums
    prose='\n'.join(p.text for p in pre if p.style.name!='Caption')
    for i in range(1,17):assert re.search(fr'图\s*{i}(?!\d)',prose),i
    for i in range(1,13):assert re.search(fr'表\s*{i}(?!\d)',prose),i
    # Compare appendix paragraphs, including full code, with the V3 Word.
    def appendix(doc):
        texts=[p.text for p in doc.paragraphs];start=next(i for i,t in enumerate(texts) if t.startswith('附录 A'));return texts[start:]
    assert appendix(old)==appendix(new),'Changed inherited appendix/code paragraphs'
    q1=load('artifacts/v3/q1_summary.json')
    q2={x['strategy']:x for x in load('artifacts/q12/q2_summary.json')['strategies']}
    q34={x['strategy']:x for x in load('artifacts/v3/summary.json')}
    evidence=[]
    def claim(label,value,fmt,source):
        token=format(value,fmt);assert token in abstract,(label,token)
        evidence.append(dict(claim=label,value=value,display=token,source=source))
    claim('Q1 purchase',q1['plan_kwh'],',.2f','artifacts/v3/q1_summary.json#plan_kwh')
    claim('Q1 cost',q1['cost'],',.2f','artifacts/v3/q1_summary.json#cost')
    claim('Q1 saving',100*(1-q1['cost']/q1['no_storage_cost']),'.2f','artifacts/v3/q1_summary.json')
    c2=q2['ridge_q0.7']['total_cost_yuan']
    claim('Q2 cost',c2,',.2f','artifacts/q12/q2_summary.json#ridge_q0.7')
    for k in ['seasonal_q0','ridge_q0']:
        claim('Q2 saving vs '+k,100*(1-c2/q2[k]['total_cost_yuan']),'.2f','artifacts/q12/q2_summary.json')
    c3=q34['q3_m7']['total_cost_yuan'];base=q34['q3_m0']['total_cost_yuan']
    claim('Q3 cost',c3,',.2f','artifacts/v3/summary.json#q3_m7')
    claim('Q3 saving yuan',base-c3,',.2f','artifacts/v3/summary.json#q3_m0,q3_m7')
    claim('Q3 saving percent',100*(1-c3/base),'.2f','artifacts/v3/summary.json#q3_m0,q3_m7')
    for k in ['q4_2_ridge','q4_3_seasonal','q4_3_gru_mean']:
        claim(k+' cost',q34[k]['total_cost_yuan'],',.2f','artifacts/v3/summary.json#'+k)
    claim('GRU saving',q34['q4_3_seasonal']['total_cost_yuan']-q34['q4_3_gru_mean']['total_cost_yuan'],',.2f','artifacts/v3/summary.json')
    errors=pd.read_csv(ROOT/'artifacts/v2/price_errors_by_issue.csv')
    maes={k:(g.mae*g.n_pairs).sum()/g.n_pairs.sum() for k,g in errors.groupby('model')}
    claim('GRU price MAE',maes['gru_mean'],'.5f','artifacts/v2/price_errors_by_issue.csv')
    claim('GRU MAE decrease vs ridge',100*(1-maes['gru_mean']/maes['ridge']),'.2f','artifacts/v2/price_errors_by_issue.csv')
    bound=pd.read_csv(ROOT/'artifacts/v3/nominal_price_bound.csv');g=bound[bound.model=='gru_mean']
    claim('GRU nominal regret',g.regret.sum(),',.2f','artifacts/v3/nominal_price_bound.csv')
    claim('GRU nominal relative regret',100*g.regret.sum()/g.nominal_cost.sum(),'.3f','artifacts/v3/nominal_price_bound.csv')
    pdf=fitz.open(ROOT/'reports/完整论文_V4.pdf');pt=[p.get_text() for p in pdf]
    app=next(i+1 for i,t in enumerate(pt) if re.search(r'^附录\s*A\s+四问',t,re.M))
    assert '关键词' in pt[0] and '问题重述' in pt[1] and app-2<=30
    assert 2<app<=len(pdf),(len(pdf),app)
    assert all(len(t.strip())>10 for t in pt)
    render=load('artifacts/v4/render_audit.json');assert render['docx_sha256']==sha(ROOT/'reports/完整论文_V4.docx') and render['pdf_sha256']==sha(ROOT/'reports/完整论文_V4.pdf')
    # Body-only diagnostic copy: exclude appendix code from prose linting, retain all main tables.
    body=Document(ROOT/'reports/完整论文_V4.docx');cut=False
    for element in list(body._element.body):
        t=''.join(element.itertext())
        if '附录 A 四问指定日期完整结果' in t:cut=True
        if cut and element.tag!=qn('w:sectPr'):body._element.body.remove(element)
    body.save(ROOT/'.qa/v4/body_only.docx')
    result=dict(status='PASS',abstract_numeric_claims=evidence,word_tables_preserved=62,display_equations_preserved=33,native_math_objects_preserved=45,inherited_images_preserved=15,new_overview_images=1,appendix_and_code_paragraphs_preserved=True,abstract_methods_bold=methods,figure_numbering_and_references='1–16, continuous and cited',main_table_references='1–12, all cited',abstract_characters=len(re.sub(r'\s','',abstract)),pages=len(pdf),pages_before_appendix=app-1,body_and_references_pages=app-2,docx_sha256=sha(ROOT/'reports/完整论文_V4.docx'),pdf_sha256=sha(ROOT/'reports/完整论文_V4.pdf'),result_workbooks={str(p.relative_to(ROOT)):sha(p) for p in sorted((ROOT/'artifacts/v3').glob('result*.xlsx'))})
    (out/'verification.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps({k:v for k,v in result.items() if k not in ['abstract_numeric_claims','result_workbooks','abstract_methods_bold']},ensure_ascii=False))

if __name__=='__main__':main()
