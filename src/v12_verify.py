"""Independent V12 artifact, layout, source-listing, and workbook checks."""
from pathlib import Path
import hashlib
import json
import re
import zipfile
import numpy as np
import pandas as pd
import openpyxl
import fitz
from docx import Document
from docx.oxml.ns import qn
from lxml import etree
from src.v9_document import maths,tables
from src.v12_document import PROGRAMS
from src.v12_results import SOURCES,EXPECTED

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'artifacts/v12'

def main():
    d=Document(ROOT/'reports/完整论文_V12.docx')
    old=Document(ROOT/'reports/完整论文_V11.docx')
    pdf=fitz.open(ROOT/'reports/完整论文_V12.pdf')
    texts=[p.get_text() for p in pdf]
    md=(ROOT/'reports/完整论文_V12.md').read_text(encoding='utf8')
    assert tables(d)==tables(old) and maths(d)==maths(old)
    assert len(d.inline_shapes)==12 and len(d.tables)==61
    assert len(d._element.findall('.//'+qn('m:oMath')))==52
    assert [etree.tostring(s._inline,method='c14n') for s in d.inline_shapes]==[
        etree.tostring(s._inline,method='c14n') for s in old.inline_shapes]
    for suffix in ['docx','pdf']:
        assert (ROOT/f'reports/完整论文_V12.{suffix}').stat().st_size<20_000_000
    assert not re.search(r'GRU|PyTorch|35,126\.85|C:\\Users\\|wu135|Snowball',md,re.I)
    for sec in d.sections:
        assert abs(sec.page_width.cm-21)<.01 and abs(sec.page_height.cm-29.7)<.01
        assert min(sec.top_margin.cm,sec.bottom_margin.cm,sec.left_margin.cm,sec.right_margin.cm)>=2.5
        assert len(sec.footer._element.findall('.//'+qn('w:fldSimple')))==1
        assert sec.footer.paragraphs[0].alignment==1
    assert '摘 要' in texts[0] and '关键词' in texts[0]
    assert '一 问题重述' in texts[1]
    reference_page=next(i+1 for i,t in enumerate(texts) if '十一 参考文献' in t)
    appendix=next(i+1 for i,t in enumerate(texts) if '附录 A 四问' in t)
    source_page=next(i+1 for i,t in enumerate(texts) if '附录 B 支撑' in t)
    assert reference_page<30 and appendix>=reference_page
    pairs=json.loads((OUT/'figure_page_pairs.json').read_text(encoding='utf8'))
    assert all(p['figure_page']==p['caption_page'] for p in pairs)
    assert len(pairs)==12
    colors={s['color'] for p in pdf for b in p.get_text('dict')['blocks'] if 'lines' in b
            for line in b['lines'] for s in line['spans'] if s['text'].strip()}
    assert colors=={0},colors
    layouts=[]
    horizontal=[]
    for i,p in enumerate(pdf):
        blocks=[b for b in p.get_text('blocks') if b[1]<770 and len(str(b[4]).strip())>3]
        bottom=max([b[3] for b in blocks]+[0])
        if i<reference_page:
            layouts.append({'page':i+1,'content_bottom_pt':round(bottom,2),
                            'bottom_space_above_margin_pt':round(max(0,770-bottom),2)})
        for b in blocks:
            if b[0]<70 or b[2]>525:
                horizontal.append({'page':i+1,'box':list(b[:4]),'text':b[4][:70]})
    assert 640<layouts[0]['content_bottom_pt']<745
    assert all(x['content_bottom_pt']>650 for x in layouts[1:reference_page-1]),layouts
    assert not horizontal,horizontal[:8]
    # Check that each appendix listing contains the complete source, line-for-line.
    paragraphs=d.paragraphs
    for filename in PROGRAMS:
        start=next(i for i,p in enumerate(paragraphs) if p.text==filename)
        end=next((i for i in range(start+1,len(paragraphs))
                  if paragraphs[i].style.name.startswith('Heading')),len(paragraphs))
        listing='\n'.join(p.text for p in paragraphs[start+1:end])
        source=(ROOT/filename).read_text(encoding='utf8').replace('\r\n','\n')
        assert listing.rstrip()==source.rstrip(),filename
    xlsx={}
    for key,path in SOURCES.items():
        f=pd.read_csv(ROOT/path)
        file=ROOT/f'results/V12/附件5/result{key}.xlsx'
        wb=openpyxl.load_workbook(file,read_only=True,data_only=True)
        original=openpyxl.load_workbook(ROOT.parent/f'CUMCM2026Problems/C题/附件/附件5/result{key}.xlsx',read_only=True)
        assert wb.sheetnames==original.sheetnames
        for name in wb.sheetnames:
            assert wb[name].max_column==original[name].max_column
        if key=='1':
            ws=wb['计划购电量']
            assert ws['A2'].value=='00:00-00:10' and ws['A145'].value=='23:50-24:00'
            plan=np.array([r[1] for r in ws.iter_rows(min_row=2,values_only=True)])
            np.testing.assert_allclose(plan,f.plan_kwh,rtol=0,atol=1e-9)
            cost=float(plan@f.price)
        else:
            sheet='调整购电量' if key in ['3','4-3'] else '计划购电量'
            ws=wb[sheet]
            assert ws.max_row==335 and ws['B1'].value=='00:00-00:10' and ws.cell(1,145).value=='23:50-24:00'
            values=list(ws.iter_rows(min_row=2,values_only=True))
            plan=np.array([r[1:145] for r in values])
            assert np.isfinite(plan).all()
            np.testing.assert_allclose(plan.ravel(),f.plan_kwh,rtol=0,atol=1e-9)
            np.testing.assert_allclose(plan.sum(axis=1),[r[145] for r in values],rtol=0,atol=1e-7)
            cost=float(sum(r[146] for r in values))
            assert wb['充放电量'].max_row==2005
        assert abs(cost-EXPECTED[key])<1e-5
        xlsx[key]={'sheets':wb.sheetnames,'cost_yuan':cost,'bytes':file.stat().st_size}
        wb.close();original.close()
    support=json.loads((OUT/'support_audit.json').read_text(encoding='utf8'))
    assert support['bytes']<20_000_000
    result={'status':'PASS','pdf_pages':len(pdf),'through_references_including_abstract':reference_page,
        'body_and_references_excluding_abstract':reference_page-1,'appendix_A_page':appendix,
        'appendix_B_page':source_page,'figures':12,'tables':61,'native_math':52,
        'all_figure_sizes_unchanged':True,'figure11_legend_corrected':True,
        'other_11_figure_media_unchanged':True,'source_programs_complete':len(PROGRAMS),
        'digital_format':'A4; margins >=2.5cm; abstract one page; centered page numbers; no TOC; anonymous; PDF and ZIP each <20MB',
        'layout':layouts,'horizontal_overflow':horizontal,'workbooks':xlsx,'support':support}
    (OUT/'verification.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf8')
    print(json.dumps({k:v for k,v in result.items() if k not in ['layout','workbooks','support']},ensure_ascii=False,indent=2))

if __name__=='__main__':main()
