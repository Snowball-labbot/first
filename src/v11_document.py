"""Revise V10 prose and typography while freezing every drawing and equation."""
from pathlib import Path
import json, re, shutil, hashlib, zipfile
from docx import Document
from docx.shared import Pt, RGBColor
from docx.oxml.ns import qn
from lxml import etree
from src.v9_document import maths, tables
ROOT=Path(__file__).resolve().parents[1]
EDITS={
23:'本问在确定性调度中引入供需预测误差。欠购触发五倍补购，多购仍按计划付费，因此先用历史残差校准日前供需输入，再估计库存用于后续缺口的价值，分别确定日前购电量与实时放电量，并以实际总购电费用评价策略。',
28:'四问在共同物理约束下依次引入供需风险、日内修正与价格不确定性，形成图1所示的统一调度框架。其中，预测提供供需和价格输入，线性规划生成购电计划，库存价值控制确定实时放电量。',
45:None,
49:'日期下标d仅出现在下标位置，放电变量d作为决策变量使用，两者按其在表达式中的位置区分。',
55:'以母线侧收支建立能量平衡，将富余供电记为未利用电量：',
61:'第一问要求日末储电量恢复至日初水平：',
63:'以购电量、充放电量、未利用电量、储电量和充放电状态为变量，得到统一的日调度模型。后续各问通过替换供需输入与初末状态条件扩展该模型：',
67:None,
73:'低价时段的外网购电同时满足负载与充电，形成集中购电峰值；高价时段由储能放电补充供电，减少外网购电量。储能由此实现购电时段的转移。',
78:'四小时汇总块内可先充后放，各十分钟区间仍满足充放电互斥。',
82:'三种储能条件采用相同供需与日初、日末状态。效率敏感性分析对五组参数分别重新求解，基准为充放电两侧各90%、往返效率81%。',
85:'按时间顺序划分训练、验证与评价阶段：1月22—31日选择参数与执行器，2月1日至12月31日开展顺序回测。模型每七天更新，使用最近至多60个已结束日期。各时点仅使用此前可获得的数据；本次评价属于同年度开发验证，泛化能力仍需跨年度检验。',
122:None,
125:'表6汇总334天、48,096个十分钟区间的功率预测误差。相较同期预测，岭回归使负载MAE由372.26降至179.86 kW，光伏MAE由171.98降至155.28 kW，对负载预测的改善更为明显。图6给出四个指定日期的日内曲线，用于分析具体时段的预测偏差；全期预测性能仍以表6的统计结果评价。',
128:None,
130:'注：误差单位为kW；光伏统计覆盖全天，包含夜间零值。',
132:'表7列出基础策略结果。图7按预测器、风险余量与库存价值控制的引入顺序汇总总费用变化；各项节费在对应的对照条件下分析。',
140:'在相同岭回归、70%余量、费用规则和初始储电量下，历史路径价值控制使334天费用由14,132,612.16元降至13,991,392.60元，节省141,219.56元（0.999%）。其中紧急购电费减少141,170.39元，而紧急电量由197,465.10增至198,137.46 kWh。节费来自补购与放电时点的重新分配：以较低价格补足当前缺口，将库存留给更高价格的后续缺口。两方案期末储电量均为5,903.65 kWh。',
144:'每次更新均以当时真实储电量为起点，固定已执行的购电与充放电量，重新优化剩余区间，构成模型预测控制的滚动决策过程[4]。',
145:'图8给出指定日期的实际光伏及四版预报，各版预报从对应发布时间起用于后续调度。',
153:'例如，原计划100 kWh、电价1元/kWh时，调整为80 kWh的费用为80＋0.5×20＝90元，调整为120 kWh的费用为100＋1.5×20＝130元。多次更新均按最终有效计划与午夜计划的净偏差结算，不累计尚未交付区间的中间调整量。',
159:'图9汇总一月验证期24组方案的费用。50%风险余量与三次日内更新的组合被选定为后续评价方案。',
163:'每次预报发布后，以当前储电量为起点优化剩余时段，仅执行至下一次更新，随后根据新信息重新求解。固定50%风险余量与基础执行器，比较三次日内更新和仅午夜预报的运行费用，以评价新增预报信息在可调单条件下的经济价值。',
174:'固定启用负载修正在1月22—31日使第三、四问滚动验证费分别增加972.85元和1,149.73元。因此采用逐月选择规则：2月使用一月选定的基础预测；3月起，每月首日分别回放此前28天的基础与修正方案。两者从窗口起点的实际储电量出发，以历史实际费用选择当月方案，费用相同时保留基础方案。候选模型、窗口、收缩系数和费用定义在全年评价前固定，月底数据仅用于次月选择。',
175:'两类电价情景均在3、6、7、8、9、10月启用修正，其余月份使用基础预测。实际储电量连续跨月传递，历史回放的末状态不覆盖实时状态。全年固定启用修正仅作消融分析，不用于事后改变选择规则。',
180:'两执行方案共享第8.5节的负载选择结果，该结果由基础执行器的历史回放生成；两者的实际库存分别连续传递，从而在相同预测规则下比较库存控制的贡献。',
190:'注：月度范围为当月价格的10%—90%分位区间，表示价格分布。',
191:'以过去28个完整日、相同发布时刻和目标时段的价格残差估计预测风险。图11给出四个指定日期的午夜预测；区间由点预测加历史残差的10%—90%分位数得到，下界不低于0.001元/kWh。全期48,096个午夜预测目标的实际覆盖率为74.36%，该范围用于描述预测误差，不作为未来价格的保证边界。',
194:None,
201:'每次滚动使用当前真实储电量、新预报和当前价格预测，已执行结果与午夜原合同保持不变。计划层以预测价格近似未来价格，按历史验证确定输入规则；供需误差由风险余量处理，价格残差区间不作为优化的硬约束。',
204:'价格变化同时影响购电时段、日内调单与库存保留。图12给出2—12月滚动策略的储能运行分布；指定日期的价格预测见图11，购电与储能结果见附录A。',
207:'注：统计范围为全期48,096个区间中的42,989个活跃区间（绝对功率大于0.01 kW）；正功率表示充电，负功率表示放电，横轴为实际交易价格。',
209:'两类策略均在2025年2月1日至12月31日逐时执行，实际储电量连续跨日传递。表10按保留合同、取消、新增和紧急补购归集费用；日前策略不能主动调单，取消与新增费用均为零。',
214:'表11汇总四问的最终结果，指定日期的购电与储能结果见附录A。',
237:'电量单位为kWh，费用单位为元，普通购电与紧急购电分列。滚动问题的指定时段和全天购电量指最终有效普通计划；午夜计划及各次更新结果见随附结果工作簿。',
}
def set_font(rpr, east='宋体', western='Times New Roman', size=None):
    fonts=rpr.get_or_add_rFonts()
    for k in list(fonts.attrib):
        if 'theme' in k.lower():del fonts.attrib[k]
    for k,v in [('ascii',western),('hAnsi',western),('cs',western),('eastAsia',east)]:fonts.set(qn('w:'+k),v)
    if size is not None:rpr.get_or_add_sz().val=Pt(size)

def main():
    src=ROOT/'reports/完整论文_V10.docx';base=Document(src);d=Document(src)
    md=src.with_suffix('.md').read_text(encoding='utf8');ps=list(d.paragraphs);log=[]
    for i,new in EDITS.items():
        p=ps[i];old=p.text
        assert not p._element.findall('.//'+qn('m:oMath')) and not p._element.findall('.//'+qn('w:drawing'))
        assert md.count(old)==1,(i,old,md.count(old))
        md=md.replace(old,new or '')
        log.append(dict(source_paragraph=i,before=old,after=new))
        if new is None:p._element.getparent().remove(p._element)
        else:
            p.runs[0].text=new
            for r in p.runs[1:]:r.text=''
    # Put the short error table before the tall, dimension-locked Figure 6.
    figure6=ps[126]._element
    table_caption=ps[129]._element
    table6=table_caption.getnext()
    assert table6.tag==qn('w:tbl')
    for element in [table_caption,table6,ps[130]._element]:figure6.addprevious(element)
    match=re.search(r'(!\[四个指定日期[^\n]+\n\n图 6[^\n]+)\s*(表 6[^\n]+\n\n\|[\s\S]+?注：误差单位为kW；光伏统计覆盖全天，包含夜间零值。)',md)
    assert match
    md=md[:match.start()]+match[2]+'\n\n'+match[1]+md[match.end():]
    for s in d.styles:
        if s.type not in [1,2]:continue
        heading=s.name.startswith('Heading');title=s.name=='Title'
        set_font(s.element.get_or_add_rPr(),'黑体' if heading or title else '宋体')
        if heading or title:s.font.color.rgb=RGBColor(0,0,0)
    for name in ['Normal','Body Text']:
        d.styles[name].font.size=Pt(12)
    appendix=False;code=False
    for p in d.paragraphs:
        t=p.text.strip();f=p.paragraph_format
        if t.startswith('附录 A'):appendix=True
        if t.startswith('附录 B'):code=True
        drawing=bool(p._element.findall('.//'+qn('w:drawing')))
        math=bool(p._element.findall('.//'+qn('m:oMath')))
        heading=p.style.name.startswith('Heading');title=p.style.name=='Title'
        caption=p.style.name=='Caption'
        size=16 if title else (14 if p.style.name=='Heading 1' else 12) if heading else 10.5 if caption or t.startswith('注：') else 10.5 if re.match(r'^\[[1-5]\]',t) else 12
        for r in p.runs:
            if r._element.find(qn('w:drawing')) is not None:continue
            oldfont=r.font.name
            set_font(r._element.get_or_add_rPr(),'黑体' if heading or title else '宋体','Consolas' if code and oldfont=='Consolas' else 'Times New Roman',None if code else size)
            if heading or title:r.bold=True
        f.widow_control=True
        if heading:
            f.keep_with_next=True;f.keep_together=True;f.space_before=Pt(9);f.space_after=Pt(5)
        elif drawing:
            f.keep_with_next=True;f.keep_together=True;f.line_spacing=1
        elif caption:
            f.keep_with_next=t.startswith('表');f.keep_together=True
        elif t.startswith('注：'):
            f.first_line_indent=Pt(0);f.keep_together=True;f.space_after=Pt(4)
        elif not appendix and not math:
            f.line_spacing=1.10;f.space_after=Pt(2)
        if re.match(r'^\[[1-5]\] ',t):
            f.keep_together=True;f.keep_with_next=not t.startswith('[5] ')
    for table in d.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    for r in p.runs:set_font(r._element.get_or_add_rPr(),size=9.5)
    # Keep a statistical note with the object it qualifies, including across tables.
    for p in d.paragraphs:
        if not p.text.startswith('注：'):continue
        prev=p._element.getprevious()
        if prev is not None and prev.tag==qn('w:p'):
            from docx.text.paragraph import Paragraph
            Paragraph(prev,p._parent).paragraph_format.keep_with_next=True
        elif prev is not None and prev.tag==qn('w:tbl'):
            from docx.table import Table
            for row in Table(prev,p._parent).rows:
                for cell in row.cells:
                    for tp in cell.paragraphs:tp.paragraph_format.keep_with_next=True
    # Headers/footers use the same explicit Chinese and Western families.
    for sec in d.sections:
        for part in [sec.header,sec.footer]:
            for p in part.paragraphs:
                for r in p.runs:set_font(r._element.get_or_add_rPr())
    assert tables(d)==tables(base) and maths(d)==maths(base)
    assert [etree.tostring(x._inline,method='c14n') for x in d.inline_shapes]==[etree.tostring(x._inline,method='c14n') for x in base.inline_shapes]
    dest=ROOT/'reports/完整论文_V11.docx';d.save(dest)
    dest.with_suffix('.md').write_text(re.sub(r'\n{3,}','\n\n',md),encoding='utf8')
    shutil.copyfile(ROOT/'reports/完整论文_V10公式.tex',ROOT/'reports/完整论文_V11公式.tex')
    out=ROOT/'artifacts/v11';out.mkdir(exist_ok=True)
    (out/'editorial_changes.json').write_text(json.dumps(log,ensure_ascii=False,indent=2),encoding='utf8')
    print(f'Built V11; {len(log)} paragraph edits; 12 drawings, 61 tables, 52 equations frozen.')
if __name__=='__main__':main()
