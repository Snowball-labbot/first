"""Check V10 evidence, editorial reduction, and native Word/PDF output."""
from pathlib import Path
import hashlib,json,re,zipfile
import numpy as np
import pandas as pd
from scipy.io import loadmat
from docx import Document
from docx.oxml.ns import qn
from lxml import etree
import fitz
from PIL import Image
from src.v9_document import maths,tables
from src.v10_document import transform
from src.q12 import audit_trace
ROOT=Path(__file__).resolve().parents[1]
def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
def main():
 out=ROOT/'artifacts/v10';a=Document(ROOT/'reports/完整论文_V9.docx');b=Document(ROOT/'reports/完整论文_V10.docx')
 skip=[]
 for i in [9,13]:
  dropped=a.inline_shapes[i]._inline.getparent()
  while dropped.tag!=qn('w:p'):dropped=dropped.getparent()
  skip.extend([dropped,dropped.getnext()])
 checks={}
 checks['only_documented_paragraph_changes']=[p.text for p in b.paragraphs]==[transform(p.text) for p in a.paragraphs if p._element not in skip]
 checks['all_61_tables_unchanged']=len(b.tables)==61 and tables(a)==tables(b)
 checks['all_52_native_equations_unchanged']=len(maths(b))==52 and maths(a)==maths(b)
 md=(ROOT/'reports/完整论文_V10.md').read_text(encoding='utf8')
 checks['12_figures_sequential']=len(b.inline_shapes)==12 and [int(x) for x in re.findall(r'!\[[^]]*\]\([^)]+\)\n\n图\s*(\d+)',md)]==list(range(1,13))
 checks['removed_draft_absent']='q3_updates.png' not in md and 'GRU' not in md and '35,126.85' not in md
 checks['all_result_workbooks_frozen']=all((out/p.name).read_bytes()==p.read_bytes() for p in (ROOT/'artifacts/v8').glob('result*.xlsx'))
 q=pd.read_csv(ROOT/'artifacts/v8/q1_intervals.csv');audit_trace(q,initial=6000.)
 checks['q1_cost_35126_95']=round(float(q.price@q.plan_kwh),2)==35126.95
 mf=json.loads((out/'figure_manifest.json').read_text(encoding='utf8'))
 checks['immutable_figure_data']=mf['source_sha256']==sha(ROOT/'artifacts/v8/figure_data.mat')
 quality={}
 for f in mf['figures']:
  p=ROOT/f"figures/v10/{f['name']}.png"
  au=json.loads(p.with_name(f['name']+'_audit.json').read_text(encoding='utf8'))
  with Image.open(p) as im:dpi=im.width/(15.5/2.54)
  quality[f['name']]=dpi
  assert dpi>=500 and au['all_text_black'] and not au['design_issues'],(f['name'],au)
  assert sha(p)==f['sha256']
  for ext in ['.pdf','.svg','.fig']:assert p.with_suffix(ext).exists()
 checks['matlab_exports_and_black_text']=True
 checks['v9_requested_heights']=all(abs(mf['figures'][n-1]['height_cm']-height)<.03 for n,height in [(2,7.9),(3,11.39172),(6,16.368),(8,8.07339)])
 r=loadmat(out/'dispatch_distribution.mat')
 checks['density_normalized']=all(abs(r[k+'_density'].sum()-100)<1e-8 for k in ['dayahead','rolling'])
 checks['density_all_334_days']=int(r['rolling_n'].item())==48096 and int(r['rolling_active'].item())==42989
 for item in json.loads((out/'dispatch_distribution_audit.json').read_text(encoding='utf8')).values():assert sha(ROOT/item['source'])==item['sha256']
 pairs=json.loads((out/'figure_page_pairs.json').read_text(encoding='utf8'))
 checks['all_figures_share_caption_page']=len(pairs)==12 and all(x['figure_page']==x['caption_page'] for x in pairs)
 pdf=fitz.open(ROOT/'reports/完整论文_V10.pdf');old=fitz.open(ROOT/'reports/完整论文_V9.pdf')
 checks['abstract_preserved']=pdf[0].get_text()==old[0].get_text()
 colors={s['color'] for p in pdf for block in p.get_text('dict')['blocks'] if 'lines' in block for l in block['lines'] for s in l['spans'] if s['text'].strip()}
 checks['pdf_text_black']=colors=={0}
 checks['no_blank_pages']=all(len(p.get_text().strip())>30 for p in pdf)
 text=[p.get_text() for p in pdf]
 appendix=next(i+1 for i,t in enumerate(text) if '附录 A 四问' in t)
 checks['body_within_30_pages']=appendix-2<=30
 checks['a4']=all(abs(p.rect.width-595.3)<1 and abs(p.rect.height-841.9)<1 for p in pdf)
 with zipfile.ZipFile(ROOT/'reports/完整论文_V10.docx') as z:
  for n in z.namelist():
   if n.endswith(('.xml','.rels')):etree.fromstring(z.read(n))
 checks['ooxml_well_formed']=True
 report=dict(passed=all(checks.values()),checks=checks,pages=len(pdf),body_and_references=appendix-2,appendix_starts=appendix,figures=12,tables=61,native_equations=52,effective_dpi=quality,
  hashes={ext:sha(ROOT/f'reports/完整论文_V10.{ext}') for ext in ['md','docx','pdf']})
 (out/'verification.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf8')
 print(json.dumps(report,ensure_ascii=False));assert report['passed']
if __name__=='__main__':main()
