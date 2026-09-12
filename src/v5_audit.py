"""Independent workbook, cumulative-inventory LP, dual certificate and trace audit."""
from pathlib import Path
import argparse,hashlib,json
import numpy as np
import pandas as pd
import openpyxl
from scipy.optimize import linprog
from src.q12 import B,forecasts,optimize_day,dump_json,audit_trace
from src.q34_data import read_extended
from src.v3_model import solve
from src.v5_model import adaptive_load

OUT=Path('artifacts/v5')
def independent_lp(load,pv,price,initial,terminal=6000,old=None,equal=False):
    """Eliminate state variables; revision uses a two-line convex epigraph."""
    n=len(price);rev=old is not None;k=(5 if rev else 4)*n;I=np.eye(n);Z=np.zeros((n,n));T=np.tril(np.ones((n,n)))
    eq=np.hstack([I,-I,I,-I]+([Z] if rev else []));rhs=np.asarray(load)-pv
    soc=np.hstack([Z,B.eta_c*T,-T/B.eta_d,Z]+([Z] if rev else []))
    ub=[soc,-soc];br=[np.full(n,B.maximum-initial),np.full(n,initial-B.minimum)]
    if equal:eq=np.vstack([eq,soc[-1]]);rhs=np.r_[rhs,terminal-initial]
    else:ub.append(-soc[-1:]);br.append(np.array([initial-terminal]))
    c=np.zeros(k);bounds=[(0,None)]*k;bounds[n:3*n]=[(0,B.limit)]*(2*n)
    if rev:
        c[4*n:]=1;bounds[4*n:]=[(None,None)]*n
        for slope in [.5,1.5]:
            row=np.zeros((n,k));row[:,:n]=np.diag(slope*price);row[:,4*n:]=-I
            ub.append(row);br.append(slope*price*old)
    else:c[:n]=price
    A=np.vstack(ub);b=np.concatenate(br)
    fit=linprog(c,A_ub=A,b_ub=b,A_eq=eq,b_eq=rhs,bounds=bounds,method='highs')
    assert fit.success,fit.message
    stationarity=c-eq.T@fit.eqlin.marginals-A.T@fit.ineqlin.marginals-fit.lower.marginals-fit.upper.marginals
    dual=float(rhs@fit.eqlin.marginals+b@fit.ineqlin.marginals)
    comp=[np.max(abs(fit.ineqlin.residual*fit.ineqlin.marginals))]
    for label,idx in [('lower',0),('upper',1)]:
        valid=np.array([v[idx] is not None for v in bounds]);v=np.array([x[idx] or 0 for x in bounds]);m=getattr(fit,label).marginals
        dual+=float(v[valid]@m[valid]);comp.append(float(np.max(abs((fit.x[valid]-v[valid])*m[valid]))) if valid.any() else 0)
    sign_violation=float(max(0,fit.ineqlin.marginals.max(),-fit.lower.marginals.min(),fit.upper.marginals.max()))
    cert=dict(objective=float(fit.fun),dual_objective=dual,gap=abs(float(fit.fun)-dual),stationarity=float(abs(stationarity).max()),dual_sign_violation=sign_violation,
        complementarity=max(comp),equality_residual=float(abs(eq@fit.x-rhs).max()),inequality_violation=float(max(0,np.max(A@fit.x-b))))
    assert max(cert[x] for x in ['gap','stationarity','complementarity','equality_residual','inequality_violation','dual_sign_violation'])<1e-5
    return cert

def main():
    ap=argparse.ArgumentParser();ap.add_argument('--data-root',required=True);a=ap.parse_args();root=Path(a.data_root)
    OUT.mkdir(parents=True,exist_ok=True);data,source=read_extended(root);books=[]
    for index in range(1,5):
        path=root/'附件'/f'附件{index}.xlsx';wb=openpyxl.load_workbook(path,data_only=False)
        sheets=[]
        for ws in wb:
            cells=list(ws.values);formula=sum(isinstance(v,str) and v.startswith('=') for row in cells for v in row)
            sheets.append(dict(name=ws.title,rows=ws.max_row,columns=ws.max_column,state=ws.sheet_state,
                formulas=formula,hidden_rows=sum(bool(v.hidden) for v in ws.row_dimensions.values()),hidden_columns=sum(bool(v.hidden) for v in ws.column_dimensions.values()),
                merges=[str(v) for v in ws.merged_cells.ranges]))
        books.append(dict(file=path.name,sha256=hashlib.sha256(path.read_bytes()).hexdigest(),sheets=sheets));wb.close()
    # Independent openpyxl extraction compared against the production pandas parser.
    wb=openpyxl.load_workbook(root/'附件/附件2.xlsx',data_only=True,read_only=True)
    for name,ws in zip(['load','pv'],wb):
        vals=np.array([r[1:] for r in list(ws.values)[1:]],float)/6
        np.testing.assert_allclose(vals,data[name],atol=3e-13,rtol=0)
    wb.close()
    pred,_=forecasts(data,'ridge',(1.,1.));ll,records=adaptive_load(data,pred['load']);changed=dict(data)
    altered=data['load'].copy();altered[20,72:]+=12345;altered[21:]+=98765;changed['load']=altered
    ll2,_=adaptive_load(changed,pred['load'])
    np.testing.assert_array_equal(ll[:20],ll2[:20]);np.testing.assert_array_equal(ll[20,:3],ll2[20,:3])
    certificates=[];rng=np.random.default_rng(20260912)
    # Randomized and adverse cases, including low prices and redundant PV.
    for case in range(36):
        n=[6,24,144][case%3];L=rng.uniform(0,800,n);V=rng.uniform(0,1000,n);p=rng.uniform(.001,2,n)
        s0=6000.;old=rng.uniform(0,900,n) if case%2 else None
        ref=independent_lp(L,V,p,s0,old=old);f,m=solve(L,V,p,s0,old=old)
        assert abs(ref['objective']-m['objective'])<1e-5
        ref.update(case=f'random_{case}',difference=abs(ref['objective']-m['objective']));certificates.append(ref)
        if old is None:
            _,mm=optimize_day(L,V,p,s0,integer=True);assert abs(mm['cost']-m['objective'])<1e-5
    L,V,p=[np.roll(data[k],1) for k in ['q1_load','q1_pv','price']]
    cert=independent_lp(L,V,p,6000,equal=True);f,m=solve(L,V,p,6000,equal=True)
    cert.update(case='q1',difference=abs(cert['objective']-m['objective']));certificates.append(cert)
    # Actual published revision instances, not only synthetic examples.
    for name in ['q3_baseline','q4_baseline','q3_adaptive','q4_adaptive','q3_online','q4_online']:
        versions=pd.read_csv(OUT/f'{name}_versions.csv.gz');trace=pd.read_csv(OUT/f'{name}_intervals.csv.gz')
        for date in ['2025-03-20','2025-06-21','2025-09-23','2025-12-21']:
            for ver in [1,2,3]:
                v=versions[(versions.date==date)&(versions.version==ver)];g=trace[(trace.date==date)&(trace.slot>=ver*36)]
                L=(v.forecast_load_kwh+v.reserve_kwh).to_numpy();V=v.forecast_pv_kwh.to_numpy();p=v.forecast_price.to_numpy();old=g.original_plan_kwh.to_numpy();state=float(v.decision_soc_kwh.iloc[0])
                cert=independent_lp(L,V,p,state,old=old);_,m=solve(L,V,p,state,old=old)
                cert.update(case=f'{name}_{date}_{ver}',difference=abs(cert['objective']-m['objective']));assert cert['difference']<1e-5;certificates.append(cert)
    pd.DataFrame(certificates).to_csv(OUT/'solver_certificates.csv',index=False)
    traces=[]
    for name in ['q3_baseline','q4_baseline','q3_adaptive','q4_adaptive','q3_online','q4_online']:
        f=pd.read_csv(OUT/f'{name}_intervals.csv.gz');d=pd.read_csv(OUT/f'{name}_daily.csv');v=pd.read_csv(OUT/f'{name}_versions.csv.gz')
        audit=audit_trace(f,initial=9902.811287870369)
        for col,key in [('load_kwh','load'),('pv_kwh','pv')]:np.testing.assert_allclose(f[col],data[key][31:].reshape(-1),atol=1e-10,rtol=0)
        expected_price=np.tile(data['price'],334) if name.startswith('q3') else data['actual_price'][31:].reshape(-1)
        np.testing.assert_allclose(f.price,expected_price,atol=1e-12,rtol=0)
        independent_cost=f.price*(np.minimum(f.original_plan_kwh,f.plan_kwh)+.5*np.maximum(f.original_plan_kwh-f.plan_kwh,0)+1.5*np.maximum(f.plan_kwh-f.original_plan_kwh,0)+5*f.emergency_kwh)
        np.testing.assert_allclose(independent_cost,f.total_cost_yuan,atol=1e-7,rtol=0)
        np.testing.assert_allclose(independent_cost.groupby(f.date).sum(),d.total_cost_yuan,atol=1e-6,rtol=0)
        for date,g in f.groupby('date',sort=False):
            plan=g.original_plan_kwh.to_numpy().copy()
            for ver,vg in v[v.date==date].groupby('version',sort=True):
                slots=vg.slot.to_numpy(int);np.testing.assert_array_equal(slots,np.arange(int(ver)*36,144));np.testing.assert_allclose(vg.old_kwh,plan[slots],atol=1e-7)
                np.testing.assert_allclose(vg.decision_soc_kwh,g.soc_start_kwh.iloc[int(ver)*36],atol=1e-7);plan[slots]=vg.new_kwh
            np.testing.assert_allclose(plan,g.plan_kwh,atol=1e-7)
        traces.append(dict(strategy=name,intervals=len(f),versions=len(v),cost=float(independent_cost.sum()),**audit))
    for new,old in [('q3_baseline','q3_m7'),('q4_baseline','q4_3_seasonal')]:
        f=pd.read_csv(OUT/f'{new}_daily.csv');g=pd.read_csv(f'artifacts/v3/{old}_daily.csv')
        np.testing.assert_allclose(f.total_cost_yuan,g.total_cost_yuan,atol=1e-6,rtol=0)
    q2=pd.read_csv('artifacts/q12/ridge_q0.7_intervals.csv.gz')
    q2_physical=audit_trace(q2,initial=9902.811287870369)
    for col,key in [('load_kwh','load'),('pv_kwh','pv')]:np.testing.assert_allclose(q2[col],data[key][31:].reshape(-1),atol=1e-10,rtol=0)
    np.testing.assert_allclose(q2.price,np.tile(data['price'],334),atol=1e-12,rtol=0)
    q2_cost=float((q2.price*(q2.plan_kwh+5*q2.emergency_kwh)).sum());assert abs(q2_cost-14132612.15924625)<1e-5
    q1=pd.read_csv('artifacts/v3/q1_intervals.csv');q1_physical=audit_trace(q1,initial=6000)
    for col,key in [('load_kwh','q1_load'),('pv_kwh','q1_pv'),('price','price')]:np.testing.assert_allclose(q1[col],np.roll(data[key],1),atol=1e-10,rtol=0)
    assert abs(float(q1.price@q1.plan_kwh)-certificates[36]['objective'])<1e-5
    choices=pd.read_csv(OUT/'online_selection.csv')
    for (mode,month),g in choices.groupby(['mode','month']):
        start=int(np.flatnonzero(data['dates'].month==month)[0]);assert (g.validation_last_day==start-1).all() and (g.validation_first_day==start-28).all()
        live=pd.read_csv(OUT/f'q{mode}_online_daily.csv');assert abs(g.initial_soc.iloc[0]-live.soc_start_kwh.iloc[start-28-31])<1e-7
        assert g[g.selected].cost.iloc[0]==g.cost.min()
    result=dict(status='PASS',workbooks=books,input_audit=source,independent_attachment2_values=105120,
        future_perturbation_test='PASS: future loads do not change earlier releases',solver_cases=len(certificates),
        q1_retained_trace=q1_physical,q2_retained_trace=q2_physical,q2_independent_cost=q2_cost,
        monthly_selection_windows=20,milp_comparisons=18,max_solver_difference=max(x['difference'] for x in certificates),
        max_duality_gap=max(x['gap'] for x in certificates),traces=traces,
        limitations=['Structural audit does not prove that sensors are error-free','Reference Modex is a supplied comparison manuscript, not independently certified results'])
    dump_json(OUT/'verification.json',result);print(json.dumps({k:result[k] for k in ['status','solver_cases','milp_comparisons','max_solver_difference','max_duality_gap']},ensure_ascii=False))
if __name__=='__main__':main()
